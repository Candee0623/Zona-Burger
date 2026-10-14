import hashlib
import logging
import re
import urllib.parse
import unicodedata
from datetime import datetime, timedelta
from decimal import Decimal, InvalidOperation
from io import BytesIO
from urllib.parse import quote
from xml.sax.saxutils import escape
from django import forms
from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.contrib.auth.hashers import make_password
from django.core.exceptions import ValidationError
from django.core.paginator import Paginator
from django.core.validators import validate_email
from django.db import IntegrityError, transaction
from django.db.models import (
    Avg,
    Count,
    DecimalField,
    ExpressionWrapper,
    F,
    Prefetch,
    Q,
    Sum,
)
from django.db.models.deletion import ProtectedError
from django.db.models.functions import TruncDate
from django.http import HttpResponse, JsonResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse
from django.utils import timezone
from django.utils.http import url_has_allowed_host_and_scheme
from django.utils.dateparse import parse_date
from django.views.decorators.http import require_POST
try:
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, LineChart, PieChart, Reference
    from openpyxl.chart.label import DataLabelList
    from openpyxl.chart.marker import DataPoint, Marker
    from openpyxl.chart.shapes import GraphicalProperties
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.properties import PageSetupProperties

    OPENPYXL_DISPONIBLE = True
except ImportError:
    OPENPYXL_DISPONIBLE = False
try:
    from reportlab.graphics.charts.barcharts import (
        HorizontalBarChart,
        VerticalBarChart,
    )
    from reportlab.graphics.charts.linecharts import HorizontalLineChart
    from reportlab.graphics.charts.piecharts import Pie
    from reportlab.graphics.shapes import Drawing, String
    from reportlab.graphics.widgets.markers import makeMarker
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_RIGHT
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import (
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    REPORTLAB_DISPONIBLE = True
except ImportError:
    REPORTLAB_DISPONIBLE = False
from .forms import ExtraForm
from .models import (
    AjusteStock, CategoriaProducto, Cliente, Compras, DetalleCompra,
    DetallePedido, Direccion, Empleado, EstadoNegocio, EstadoPedido, Extras, GrupoOpcion,
    Horario, Insumo, Lote, MedioPago, MovimientoStock, Negocio, Opcion, Pedido, Permiso,
    Producto, ProductoExtras, ProductoGrupoOpcion, ProductoPromocion,
    Promocion, Receta, RecetaInsumo, RedesSociales, Rol, RolPermiso,
    TipoBeneficio, UltimaModificacion, ZonasEntrega,
)
from .stock_service import (
    calcular_consumo_carrito,
    descontar_stock_pedido,
    reintegrar_stock_pedido,
    verificar_stock_disponible,
    StockInsuficienteError,
)

logger = logging.getLogger(__name__)

# IDs de la tabla EstadoProducto.
ESTADO_ACTIVO = 1
ESTADO_OCULTO = 2
ESTADO_ELIMINADO = 3

# Estados que no se muestran en el menú público.
ESTADOS_NO_VISIBLES = [ESTADO_OCULTO, ESTADO_ELIMINADO]

# Formato de WhatsApp para Argentina: 549 + código de área + número.
PATRON_WHATSAPP_AR = re.compile(r'^549\d{9,10}$')


# ============================================================
# UTILIDADES
# ============================================================

def limpiar_numero_telefono(numero):
    return re.sub(r'\D', '', numero or '')


def es_numero_whatsapp_valido(numero_limpio):
    return bool(PATRON_WHATSAPP_AR.fullmatch(numero_limpio))


def es_telefono_cliente_valido(numero_limpio):
    return bool(re.fullmatch(r'\d{8,13}', numero_limpio))


def obtener_instagram():
    return RedesSociales.objects.filter(
        plataforma__icontains='instagram'
    ).first()


def es_peticion_ajax(request):
    return (
        request.headers.get('x-requested-with') == 'XMLHttpRequest'
        or 'application/json' in request.headers.get('Accept', '')
    )


def calcular_ingresos_del_dia(pedidos):
    """Devuelve (efectivo, transferencia) cobrado, sin pedidos cancelados."""
    validos = pedidos.filter(pagado=True).exclude(
        idestadopedido__descripcion__iexact='Cancelado'
    )

    efectivo = validos.filter(
        idmediopago__nombremetodo__iexact='efectivo'
    ).aggregate(suma=Sum('total'))['suma'] or 0

    transferencia = validos.filter(
        idmediopago__nombremetodo__icontains='transferencia'
    ).aggregate(suma=Sum('total'))['suma'] or 0

    return efectivo, transferencia


# ============================================================
# 1. TIENDA PÚBLICA
# ============================================================

def index(request):
    negocio = Negocio.objects.first()

    horarios = Horario.objects.filter(idnegocio=negocio) if negocio else []
    zonas = (
        ZonasEntrega.objects.filter(idnegocio=negocio.idnegocio)
        if negocio else []
    )

    return render(request, 'publico/index.html', {
        'negocio': negocio,
        'horarios': horarios,
        'zonas': zonas,
        'instagram': obtener_instagram(),
        'promociones_activas': Promocion.objects.filter(activo=1),
    })


def menu(request):
    productos_visibles = Producto.objects.exclude(
        idestadoproducto_id__in=ESTADOS_NO_VISIBLES
    )

    categorias = CategoriaProducto.objects.prefetch_related(
        Prefetch('producto_set', queryset=productos_visibles)
    )

    return render(request, 'publico/menu.html', {
        'categorias': categorias,
        'negocio': Negocio.objects.first(),
        'instagram': obtener_instagram(),
    })


def menu_actualizaciones(request):
    """
    Permite que el menú público se actualice sin recargar la página.

    El navegador envía la versión que tiene en el header X-Menu-Version.
    Si coincide con la actual responde {"cambios": false}; si no, devuelve
    los fragmentos HTML actualizados de productos y categorías.
    """
    categorias = CategoriaProducto.objects.prefetch_related(
        Prefetch(
            'producto_set',
            queryset=(
                Producto.objects
                .select_related('idestadoproducto')
                .exclude(idestadoproducto_id__in=ESTADOS_NO_VISIBLES)
            ),
        )
    )

    # Si algo visible cambia (nombre, precio, imagen, estado, categoría...)
    # cambia el hash resultante.
    datos_menu = []

    for categoria in categorias:
        productos = []

        for producto in categoria.producto_set.all():
            productos.append({
                'id': producto.idproducto,
                'nombre': producto.nombre,
                'descripcion': producto.descripcion or '',
                'precio': str(producto.precio),
                'imagen': producto.imagen.name if producto.imagen else '',
                'estado': (
                    producto.idestadoproducto.descripcion
                    if producto.idestadoproducto else ''
                ),
            })

        datos_menu.append({
            'id': categoria.pk,
            'nombre': categoria.nombre,
            'productos': productos,
        })

    # MD5 solo para detectar cambios, no se usa con fines de seguridad.
    version_actual = hashlib.md5(
        repr(datos_menu).encode('utf-8')
    ).hexdigest()

    if request.headers.get('X-Menu-Version') == version_actual:
        return JsonResponse({'cambios': False, 'version': version_actual})

    html_productos = render(
        request,
        'components/menuProductos.html',
        {'categorias': categorias},
    ).content.decode('utf-8')

    html_categorias = render(
        request,
        'components/menuCategorias.html',
        {'categorias': categorias},
    ).content.decode('utf-8')

    return JsonResponse({
        'cambios': True,
        'version': version_actual,
        'menu': html_productos,
        'categorias': html_categorias,
    })


def detalleProducto(request, idproducto):
    producto = get_object_or_404(Producto, pk=idproducto)

    relaciones_grupos = (
        ProductoGrupoOpcion.objects
        .filter(idproducto=producto)
        .select_related('idgrupo')
    )

    grupos_opciones = [
        {'grupo': rel.idgrupo, 'opciones': rel.idgrupo.opcion_set.all()}
        for rel in relaciones_grupos
        if rel.idgrupo
    ]

    extras_producto = [
        rel.idextra
        for rel in producto.productoextras_set.select_related('idextra')
        if rel.idextra
    ]

    # Ingredientes de la receta que el cliente puede quitar.
    insumos_removibles = (
        RecetaInsumo.objects
        .filter(
            idreceta__idproducto=producto,
            es_removible=True,
            idinsumo__isnull=False,
        )
        .select_related('idinsumo')
    )

    # Si el cliente está editando un ítem, se recupera su configuración.
    edit_item_id = request.session.get('edit_item_id')
    carrito = request.session.get('carrito', {})
    item_editando = carrito.get(edit_item_id) if edit_item_id else None

    return render(request, 'publico/detalleProducto.html', {
        'producto': producto,
        'grupos_opciones': grupos_opciones,
        'extras_producto': extras_producto,
        'insumos_removibles': insumos_removibles,
        'negocio': Negocio.objects.first(),
        'instagram': obtener_instagram(),
        'item_editando': item_editando,
        'edit_item_id': edit_item_id,
    })


# ============================================================
# 2. CARRITO
# ============================================================

def agregar_al_carrito(request, idproducto):
    if request.method != 'POST':
        return redirect('menu')

    producto = get_object_or_404(Producto, pk=idproducto)

    # Cada grupo de opciones puede exigir una cantidad mínima de selecciones.
    relaciones_grupos = (
        ProductoGrupoOpcion.objects
        .filter(idproducto=producto)
        .select_related('idgrupo')
    )

    for relacion in relaciones_grupos:
        grupo = relacion.idgrupo

        if not grupo:
            continue

        min_selecciones = int(grupo.minselecciones or 0)

        if min_selecciones <= 0:
            continue

        seleccionadas = request.POST.getlist(f'grupo_{grupo.idgrupo}')

        if len(seleccionadas) < min_selecciones:
            # Se conserva el ítem en edición para volver a cargarlo.
            edit_id = request.POST.get('edit_id')

            if edit_id:
                request.session['edit_item_id'] = edit_id

            messages.error(
                request,
                f'Debés seleccionar al menos {min_selecciones} '
                f'opción(es) en "{grupo.nombre}".'
            )

            return redirect('detalleProducto', idproducto=producto.idproducto)

    precio_unitario = float(producto.precio)
    detalle_opciones = []
    opciones_seleccionadas = []
    extras_seleccionados = {}
    insumos_quitados_ids = []

    for key in request.POST:
        if key.startswith('grupo_'):
            for value in request.POST.getlist(key):
                opcion = Opcion.objects.filter(pk=value).first()

                if opcion:
                    precio_unitario += float(opcion.precioadicional or 0)
                    detalle_opciones.append(opcion.nombre)
                    opciones_seleccionadas.append(str(value))

    for key, value in request.POST.items():
        if key.startswith('extra_'):
            extra = Extras.objects.filter(pk=value).first()

            if extra:
                cantidad_extra = int(
                    request.POST.get(f'cantidad_extra_{value}', 1)
                )
                precio_unitario += float(extra.precio or 0) * cantidad_extra
                detalle_opciones.append(f'{extra.nombre} (x{cantidad_extra})')
                extras_seleccionados[str(value)] = cantidad_extra

        elif key.startswith('sin_insumo_'):
            insumo = Insumo.objects.filter(pk=value).first()

            if insumo:
                detalle_opciones.append(f'Sin {insumo.nombreinsumo}')
                insumos_quitados_ids.append(str(value))

    carrito = request.session.setdefault('carrito', {})

    # Al editar un ítem se reemplaza la versión anterior.
    edit_id = request.POST.get('edit_id')

    if edit_id and edit_id in carrito:
        del carrito[edit_id]

    request.session.pop('edit_item_id', None)

    # El ID combina producto y opciones: dos configuraciones distintas del
    # mismo producto quedan como ítems separados.
    item_id = (
        f"{idproducto}_" + '_'.join(detalle_opciones)
        if detalle_opciones else str(idproducto)
    )

    if item_id in carrito:
        carrito[item_id]['cantidad'] += 1
        carrito[item_id]['subtotal'] = (
            carrito[item_id]['cantidad'] * precio_unitario
        )
    else:
        carrito[item_id] = {
            'producto_id': producto.idproducto,
            'nombre': producto.nombre,
            'imagen': producto.imagen.url if producto.imagen else '',
            'precio_unitario': precio_unitario,
            'cantidad': 1,
            'subtotal': precio_unitario,
            'detalles': detalle_opciones,
            'opciones_ids': opciones_seleccionadas,
            'extras_ids': extras_seleccionados,
            'insumos_quitados_ids': insumos_quitados_ids,
        }

    # Django no detecta cambios dentro de diccionarios anidados de la sesión.
    request.session.modified = True

    return redirect('ver_carrito')


def editar_item_carrito(request, item_id):
    carrito = request.session.get('carrito', {})

    if item_id in carrito:
        request.session['edit_item_id'] = item_id
        return redirect(
            'detalleProducto',
            idproducto=carrito[item_id]['producto_id'],
        )

    return redirect('ver_carrito')


def ver_carrito(request):
    # Al volver al carrito se descarta el estado de edición.
    if request.session.pop('edit_item_id', None) is not None:
        request.session.modified = True

    carrito = request.session.get('carrito', {})
    total_carrito = sum(item['subtotal'] for item in carrito.values())

    return render(request, 'publico/carrito.html', {
        'carrito': carrito,
        'total_carrito': total_carrito,
        'negocio': Negocio.objects.first(),
        'instagram': obtener_instagram(),
    })


def actualizar_cantidad_carrito(request, item_id, accion):
    carrito = request.session.get('carrito', {})

    if item_id in carrito:
        if accion == 'aumentar':
            carrito[item_id]['cantidad'] += 1
        elif accion == 'disminuir':
            carrito[item_id]['cantidad'] -= 1

        if carrito[item_id]['cantidad'] <= 0:
            del carrito[item_id]
        else:
            carrito[item_id]['subtotal'] = (
                carrito[item_id]['cantidad']
                * float(carrito[item_id]['precio_unitario'])
            )

        request.session.modified = True

    return redirect('ver_carrito')


def eliminar_del_carrito(request, item_id):
    carrito = request.session.get('carrito', {})

    if item_id in carrito:
        del carrito[item_id]
        request.session.modified = True

    return redirect('ver_carrito')


# ============================================================
# 3. LÓGICA DE PROMOCIONES
# ============================================================

def promocion_es_valida(promocion):
    if not promocion:
        return False

    if promocion.activo != 1:
        return False

    hoy = timezone.localdate()

    if promocion.fechainicio and hoy < promocion.fechainicio:
        return False

    if promocion.fechafin and hoy > promocion.fechafin:
        return False

    return True


def obtener_productos_promocion(promocion):
    return set(
        ProductoPromocion.objects
        .filter(idpromocion=promocion)
        .values_list('idproducto_id', flat=True)
    )


def calcular_descuento_promocion(carrito, promocion):
    """Devuelve (descuento, mensaje) según el tipo de beneficio."""
    if not promocion:
        return Decimal('0.00'), ''

    if not promocion_es_valida(promocion):
        return Decimal('0.00'), 'La promoción no está vigente.'

    productos_validos = obtener_productos_promocion(promocion)

    if not productos_validos:
        return Decimal('0.00'), 'La promoción no tiene productos asociados.'

    # Una entrada por unidad, para poder aplicar promos tipo 2x1 o 3x2.
    unidades = []
    subtotal_elegible = Decimal('0.00')

    for item in carrito.values():
        producto_id = item.get('producto_id')

        if producto_id is None:
            continue

        try:
            producto_id = int(producto_id)
        except (TypeError, ValueError):
            continue

        if producto_id not in productos_validos:
            continue

        try:
            cantidad = int(item.get('cantidad', 0))
        except (TypeError, ValueError):
            cantidad = 0

        if cantidad <= 0:
            continue

        try:
            precio_unitario = Decimal(str(item.get('precio_unitario', 0)))
        except Exception:
            precio_unitario = Decimal('0.00')

        if precio_unitario <= 0:
            continue

        subtotal_elegible += precio_unitario * cantidad
        unidades.extend([precio_unitario] * cantidad)

    if not unidades:
        return (
            Decimal('0.00'),
            'La promoción no aplica a los productos del carrito.',
        )

    tipo = promocion.tipobeneficio.codigo if promocion.tipobeneficio else None
    valor = promocion.valor or Decimal('0.00')
    descuento = Decimal('0.00')
    mensaje = ''

    if tipo == '2x1':
        unidades.sort()
        descuento = sum(unidades[:len(unidades) // 2], Decimal('0.00'))
        mensaje = 'Promoción 2 x 1 aplicada.'

    elif tipo == '3x2':
        unidades.sort()
        descuento = sum(unidades[:len(unidades) // 3], Decimal('0.00'))
        mensaje = 'Promoción 3 x 2 aplicada.'

    elif tipo == 'PRODUCTO_GRATIS':
        descuento = min(unidades)
        mensaje = 'Producto gratis aplicado.'

    elif tipo == 'DESCUENTO_PORCENTAJE':
        porcentaje = max(Decimal('0.00'), min(valor, Decimal('100.00')))
        descuento = subtotal_elegible * porcentaje / Decimal('100')
        mensaje = f'Descuento del {porcentaje}% aplicado.'

    elif tipo == 'DESCUENTO_FIJO':
        descuento = min(valor, subtotal_elegible)
        mensaje = f'Descuento de ${descuento} aplicado.'

    else:
        return Decimal('0.00'), 'Tipo de promoción no válido.'

    descuento = max(Decimal('0.00'), min(descuento, subtotal_elegible))
    descuento = descuento.quantize(Decimal('0.01'))

    return descuento, mensaje


def calcular_total_pedido(carrito, promocion=None):
    subtotal = Decimal('0.00')

    for item in carrito.values():
        try:
            cantidad = int(item.get('cantidad', 0))
        except (TypeError, ValueError):
            cantidad = 0

        try:
            precio_unitario = Decimal(str(item.get('precio_unitario', 0)))
        except Exception:
            precio_unitario = Decimal('0.00')

        if cantidad > 0:
            subtotal += precio_unitario * cantidad

    descuento, mensaje = calcular_descuento_promocion(carrito, promocion)

    total = max(Decimal('0.00'), subtotal - descuento)

    return {
        'subtotal': subtotal.quantize(Decimal('0.01')),
        'descuento': descuento.quantize(Decimal('0.01')),
        'total': total.quantize(Decimal('0.01')),
        'mensaje': mensaje,
    }


@require_POST
def aplicar_promocion(request):
    codigo = request.POST.get('codigo_promocion', '').strip().upper()

    if not codigo:
        return JsonResponse({
            'ok': False,
            'mensaje': 'Ingresá una palabra clave.',
        })

    try:
        promocion = Promocion.objects.get(
            palabraclave__iexact=codigo,
            activo=1,
        )
    except Promocion.DoesNotExist:
        return JsonResponse({
            'ok': False,
            'mensaje': 'La promoción no existe o no está activa.',
        })

    if not promocion_es_valida(promocion):
        return JsonResponse({
            'ok': False,
            'mensaje': 'La promoción no está vigente.',
        })

    carrito = request.session.get('carrito', {})

    if not carrito:
        return JsonResponse({
            'ok': False,
            'mensaje': 'El carrito está vacío.',
        })

    resultado = calcular_total_pedido(carrito, promocion)

    if resultado['descuento'] <= Decimal('0.00'):
        return JsonResponse({
            'ok': False,
            'mensaje': (
                resultado['mensaje'] or 'La promoción no aplica al carrito.'
            ),
        })

    request.session['promocion_id'] = promocion.idpromocion
    request.session['codigo_promocion'] = codigo
    request.session['descuento_promocion'] = str(resultado['descuento'])
    request.session.modified = True

    return JsonResponse({
        'ok': True,
        'promocion_id': promocion.idpromocion,
        'codigo': codigo,
        'descuento': float(resultado['descuento']),
        'subtotal': float(resultado['subtotal']),
        'total': float(resultado['total']),
        'mensaje': resultado['mensaje'],
    })


# ============================================================
# 4. CHECKOUT Y PEDIDOS
# ============================================================

def _items_checkout(carrito):
    """Convierte los importes a Decimal y calcula el precio unitario real."""
    items = []

    for item in carrito.values():
        cantidad = item.get('cantidad', 1) or 1
        subtotal = Decimal(str(item.get('subtotal', 0)))

        items.append({
            **item,
            'subtotal': subtotal,
            'precio_unitario': subtotal / cantidad,
        })

    return items


def calcular_consumo_insumos(carrito):
    """Devuelve el consumo total de insumos del carrito (recetas base y opciones)."""
    return calcular_consumo_carrito(carrito)


def procesar_checkout(request):
    carrito = request.session.get('carrito', {})

    total_carrito = sum(
        (Decimal(str(i.get('subtotal', 0))) for i in carrito.values()),
        Decimal('0.00'),
    )

    medios_pago = list(
        MedioPago.objects.values('idmediopago', 'nombremetodo')
    )

    zonas_entrega = list(
        ZonasEntrega.objects.values('idzona', 'nombre', 'costoenvio')
    )

    es_consulta_json = (
        request.headers.get('x-requested-with') == 'XMLHttpRequest'
        or (
            request.method == 'GET'
            and 'application/json' in request.META.get('HTTP_ACCEPT', '')
        )
    )

    if es_consulta_json:
        return JsonResponse({
            'status': 'success',
            'total_carrito': str(total_carrito),
            'medios_pago': medios_pago,
            'zonas_entrega': zonas_entrega,
        })

    contexto = {
        'medios_pago': medios_pago,
        'zonas_entrega': zonas_entrega,
        'total_carrito': total_carrito,
        'carrito_items': _items_checkout(carrito),
    }

    if request.method != 'POST':
        return render(request, 'publico/checkout.html', {
            **contexto,
            'total_estimado': total_carrito,
            'errores': {'codigo_promocion': ''},
            'codigo_promocion': '',
            'datos_previos': {},
            'hay_datos_previos': False,
            'error_telefono': '',
            'error_medio_pago': '',
            'error_promocion': '',
            'descuento_promocion': Decimal('0.00'),
            'objeto_promocion': None,
        })

    def campo(nombre):
        return (request.POST.get(nombre) or '').strip()

    nombre = campo('nombre')
    apellido = campo('apellido')
    telefono = campo('telefono')
    calle = campo('calle')
    numero = campo('numero')
    piso = campo('piso')
    localidad = campo('localidad') or 'A coordinar con el vendedor'
    horario_entrega = campo('horario_entrega')
    codigo_promo = campo('codigo_promocion').upper()
    comentario = campo('comentario')
    id_medio_pago = request.POST.get('medio_pago')

    if not carrito:
        messages.error(request, 'Tu carrito está vacío.')
        return redirect('ver_carrito')

    errores = {'codigo_promocion': ''}

    if not nombre:
        errores['nombre'] = 'Ingresá tu nombre.'
    elif not 2 <= len(nombre) <= 100:
        errores['nombre'] = 'El nombre debe tener entre 2 y 100 caracteres.'

    if not apellido:
        errores['apellido'] = 'Ingresá tu apellido.'
    elif not 2 <= len(apellido) <= 100:
        errores['apellido'] = 'El apellido debe tener entre 2 y 100 caracteres.'

    telefono_limpio = limpiar_numero_telefono(telefono)

    if not es_telefono_cliente_valido(telefono_limpio):
        errores['telefono'] = (
            'Ingresá un teléfono válido '
            '(solo números, entre 8 y 13 dígitos).'
        )

    if not calle:
        errores['calle'] = 'Ingresá la calle.'
    elif len(calle) > 150:
        errores['calle'] = 'La calle es demasiado larga.'

    if not numero:
        errores['numero'] = 'Ingresá el número de la calle.'
    elif len(numero) > 20:
        errores['numero'] = 'El número de calle es demasiado largo.'

    medio_pago = None

    if not id_medio_pago:
        errores['medio_pago'] = 'Seleccioná un método de pago.'
    else:
        medio_pago = MedioPago.objects.filter(pk=id_medio_pago).first()

        if medio_pago is None:
            errores['medio_pago'] = 'Seleccioná un método de pago válido.'

    horario_deseado = None

    if horario_entrega:
        try:
            horario_deseado = datetime.strptime(
                horario_entrega, '%H:%M'
            ).time()
        except ValueError:
            errores['horario_entrega'] = 'El horario de entrega no es válido.'

    zona = ZonasEntrega.objects.filter(nombre__iexact=localidad).first()

    if zona and zona.costoenvio is not None:
        costo_envio = Decimal(str(zona.costoenvio))
        costo_envio_texto = (
            'Envío gratis' if costo_envio == 0 else f'${costo_envio:.2f}'
        )
    else:
        costo_envio = Decimal('0.00')
        costo_envio_texto = 'A coordinar'

    objeto_promocion = None
    descuento_promocion = Decimal('0.00')
    error_promocion = ''

    if codigo_promo:
        objeto_promocion = Promocion.objects.filter(
            palabraclave__iexact=codigo_promo
        ).first()

        if not objeto_promocion:
            error_promocion = 'El código promocional no existe.'

        elif not promocion_es_valida(objeto_promocion):
            error_promocion = 'La promoción no está vigente.'

        else:
            descuento_promocion, mensaje_promocion = (
                calcular_descuento_promocion(carrito, objeto_promocion)
            )

            if descuento_promocion <= 0:
                error_promocion = (
                    mensaje_promocion
                    or 'La promoción no aplica a los productos de tu carrito.'
                )

        if error_promocion:
            errores['codigo_promocion'] = error_promocion

    if any(errores.values()):
        total_estimado = max(
            Decimal('0.00'),
            total_carrito - descuento_promocion,
        )

        return render(request, 'publico/checkout.html', {
            **contexto,
            'total_estimado': total_estimado,
            'datos_previos': request.POST,
            'hay_datos_previos': True,
            'errores': errores,
            'error_telefono': errores.get('telefono', ''),
            'error_medio_pago': errores.get('medio_pago', ''),
            'error_promocion': error_promocion,
            'descuento_promocion': descuento_promocion,
            'objeto_promocion': objeto_promocion,
        })

    telefono = telefono_limpio

    total_general = max(
        Decimal('0.00'),
        total_carrito - descuento_promocion + costo_envio,
    )

    consumo_insumos = calcular_consumo_carrito(carrito)

    # 1. Verificación previa de stock disponible
    stock_ok, errores_stock = verificar_stock_disponible(consumo_insumos)
    if not stock_ok:
        for err in errores_stock:
            messages.error(request, err)
        return redirect('ver_carrito')

    try:
        with transaction.atomic():
            cliente = Cliente.objects.create(
                nombre=nombre,
                apellido=apellido,
                telefono=telefono,
            )

            Direccion.objects.create(
                idcliente=cliente,
                calle=calle,
                numero=numero,
                piso=piso,
                localidad=localidad,
            )

            estado_inicial = get_object_or_404(EstadoPedido, pk=1)

            pedido = Pedido.objects.create(
                idcliente=cliente,
                idmediopago=medio_pago,
                idestadopedido=estado_inicial,
                idpromocion=(
                    objeto_promocion.idpromocion if objeto_promocion else None
                ),
                horarioentregadeseado=horario_deseado,
                total=total_general,
                justificacioncancelacion=comentario,
            )

            for item in carrito.values():
                producto = Producto.objects.filter(
                    pk=item.get('producto_id')
                ).first()

                if not producto:
                    continue

                cantidad = item.get('cantidad', 1) or 1
                subtotal_item = Decimal(str(item.get('subtotal', 0)))
                detalles = item.get('detalles')

                DetallePedido.objects.create(
                    idpedido=pedido,
                    idproducto=producto,
                    cantidad=cantidad,
                    preciounitario=subtotal_item / cantidad,
                    observaciones=(
                        ', '.join(detalles) if isinstance(detalles, list) else ''
                    ),
                )

            # Descontar stock atómicamente y registrar movimientos
            descontar_stock_pedido(pedido, consumo_insumos)

    except StockInsuficienteError as e:
        messages.error(request, str(e))
        return redirect('ver_carrito')

    mensaje = '*¡Nuevo Pedido! 🍔*\n\n'

    mensaje += '*Datos del Cliente:*\n'
    mensaje += f'• Nombre: {nombre} {apellido}\n'
    mensaje += f'• Teléfono: {telefono}\n\n'

    mensaje += '*Dirección de Envío:*\n'
    mensaje += f'• Calle: {calle} {numero}'

    if piso:
        mensaje += f' (Piso/Depto: {piso})'

    mensaje += f'\n• Localidad / Zona: {localidad}\n'

    if horario_entrega:
        mensaje += f'• Horario Deseado: {horario_entrega}\n'

    if objeto_promocion and descuento_promocion > 0:
        mensaje += f'• Palabra Clave Promo: {codigo_promo} (Aplicada)\n'

    mensaje += '\n*Detalle del Pedido:*\n'

    for item in carrito.values():
        subtotal = Decimal(str(item.get('subtotal', 0)))

        mensaje += (
            f'• {item.get("cantidad")}x '
            f'{item.get("nombre")} '
            f'(${subtotal:.2f})\n'
        )

        for detalle in item.get('detalles', []):
            mensaje += f'   - {detalle}\n'

    mensaje += f'\n• Subtotal Productos: ${total_carrito:.2f}\n'

    if descuento_promocion > 0:
        mensaje += f'• Descuento Promo: -${descuento_promocion:.2f}\n'

    mensaje += f'• Costo de Envío: {costo_envio_texto}\n'
    mensaje += f'*Total a Pagar:* ${total_general:.2f}\n'
    mensaje += f'*Método de Pago:* {medio_pago.nombremetodo}\n'

    if comentario:
        mensaje += f'*Comentarios:* {comentario}\n'

    negocio = Negocio.objects.first()

    numero_whatsapp = limpiar_numero_telefono(
        str(negocio.telefono) if negocio and negocio.telefono else ''
    )

    if 'carrito' in request.session:
        del request.session['carrito']
        request.session.modified = True

    if not numero_whatsapp:
        logger.warning(
            'El negocio no tiene un número de WhatsApp configurado.'
        )
        return redirect('menu')

    request.session['whatsapp_url'] = (
        f'https://wa.me/{numero_whatsapp}?text={quote(mensaje)}'
    )
    request.session['limpiar_checkout'] = True
    request.session.modified = True

    return redirect('pedido_exitoso')


def pedido_exitoso(request):
    limpiar_checkout = request.session.pop('limpiar_checkout', False)

    return render(request, 'publico/pedidoExitoso.html', {
        'negocio': Negocio.objects.first(),
        'whatsapp_url': request.session.get('whatsapp_url'),
        'limpiar_checkout': limpiar_checkout,
    })


# ============================================================
# 5. PANEL - INICIO Y PEDIDOS
# ============================================================

def panel_inicio(request):
    hoy = timezone.localdate()
    pedidos_hoy_base = Pedido.objects.filter(fecha_creacion__date=hoy)

    filtro_estado = request.GET.get('estado')
    filtro_cliente = request.GET.get('cliente', '').strip()

    nombres_estados = [
        'Pendiente',
        'En Preparación',
        'En proceso',
        'Enviado',
        'Entregado',
        'Cancelado',
    ]

    estados_con_conteo = []

    for nombre in nombres_estados:
        estado = EstadoPedido.objects.filter(
            descripcion__iexact=nombre
        ).first()

        # Si el estado todavía no existe en la base, se omite del resumen.
        if estado:
            estados_con_conteo.append({
                'estado': estado,
                'cantidad': pedidos_hoy_base.filter(
                    idestadopedido=estado
                ).count(),
            })

    pedidos_hoy = pedidos_hoy_base

    if filtro_estado:
        pedidos_hoy = pedidos_hoy.filter(idestadopedido_id=filtro_estado)

    # Cada palabra debe coincidir con el comienzo del nombre o del apellido.
    if filtro_cliente:
        for termino in filtro_cliente.split():
            pedidos_hoy = pedidos_hoy.filter(
                Q(idcliente__nombre__istartswith=termino)
                | Q(idcliente__apellido__istartswith=termino)
            )

    total_efectivo, total_transferencia = calcular_ingresos_del_dia(
        pedidos_hoy_base
    )

    return render(request, 'panel/inicioPanel.html', {
        'pedidos_hoy': pedidos_hoy,
        'estados_con_conteo': estados_con_conteo,
        'filtro_estado_actual': filtro_estado,
        'filtro_cliente_actual': filtro_cliente,
        'total_pedidos_hoy': pedidos_hoy_base.count(),
        'total_efectivo': total_efectivo,
        'total_transferencia': total_transferencia,
    })


def obtener_ticket_modal(request, idpedido):
    pedido = get_object_or_404(Pedido, idpedido=idpedido)

    direccion_cliente = (
        Direccion.objects.filter(idcliente=pedido.idcliente).first()
        if pedido.idcliente else None
    )

    detalles = DetallePedido.objects.filter(idpedido=pedido)

    # El subtotal de cada producto se calcula solo para mostrarlo en el ticket.
    for detalle in detalles:
        detalle.subtotal_item = detalle.preciounitario * detalle.cantidad

    return render(request, 'panel/parcialTicket.html', {
        'pedido': pedido,
        'detalles': detalles,
        'direccion_cliente': direccion_cliente,
    })


def actualizar_estado_pedido(request, idpedido):
    if request.method != 'POST':
        return redirect('panel_inicio')

    pedido = get_object_or_404(Pedido, idpedido=idpedido)
    nuevo_estado_id = request.POST.get('nuevo_estado')

    if not nuevo_estado_id:
        return redirect('panel_inicio')

    nuevo_estado = get_object_or_404(EstadoPedido, pk=nuevo_estado_id)
    pedido.idestadopedido = nuevo_estado

    es_cancelado = 'cancelado' in nuevo_estado.descripcion.lower()

    if es_cancelado:
        pedido.justificacioncancelacion = request.POST.get(
            'justificacioncancelacion', ''
        )
        reintegrar_stock_pedido(pedido, motivo=pedido.justificacioncancelacion)

    pedido.save()

    cliente = pedido.idcliente
    nombre_cliente = cliente.nombre if cliente else 'Cliente'
    telefono_cliente = getattr(cliente, 'telefono', '') if cliente else ''

    if es_cancelado:
        motivo = pedido.justificacioncancelacion or 'Sin motivo especificado'
        mensaje_texto = (
            f'Hola {nombre_cliente}, te escribimos de Zona Burger para '
            f'informarte que tu pedido #{pedido.idpedido} ha sido '
            f'*CANCELADO*. Motivo: {motivo}.'
        )
    else:
        mensaje_texto = (
            f'Hola {nombre_cliente}, te escribimos de Zona Burger. '
            f'Tu pedido #{pedido.idpedido} ahora se encuentra en estado: '
            f'*{nuevo_estado.descripcion}*.'
        )

    messages.success(
        request,
        f'Estado actualizado correctamente. '
        f'Mensaje preparado para {nombre_cliente}.',
    )

    # El panel abre el enlace de WhatsApp a partir del parámetro wa_url.
    if telefono_cliente:
        url_whatsapp = (
            f'https://wa.me/{telefono_cliente}'
            f'?text={urllib.parse.quote(mensaje_texto)}'
        )

        return redirect(
            f"{reverse('panel_inicio')}"
            f"?wa_url={urllib.parse.quote(url_whatsapp)}"
        )

    return redirect('panel_inicio')


def toggle_pagado(request, idpedido):
    pedido = get_object_or_404(Pedido, pk=idpedido)
    pedido.pagado = not pedido.pagado
    pedido.save()

    ingresos_efectivo, ingresos_transferencia = calcular_ingresos_del_dia(
        Pedido.objects.filter(fecha_creacion__date=timezone.localdate())
    )

    if es_peticion_ajax(request):
        return JsonResponse({
            'success': True,
            'pagado': pedido.pagado,
            'total_efectivo': float(ingresos_efectivo),
            'total_transferencia': float(ingresos_transferencia),
        })

    return redirect('panel_inicio')


# Pendientes de implementar.

def ver_ticket(request, idpedido):
    get_object_or_404(Pedido, idpedido=idpedido)
    return HttpResponse(f'Visualizando el ticket del pedido #{idpedido}')


def imprimir_ticket(request, idpedido):
    get_object_or_404(Pedido, idpedido=idpedido)
    return HttpResponse(f'Imprimiendo ticket del pedido #{idpedido}')


def cambiar_estado_pedido(request, idpedido):
    get_object_or_404(Pedido, idpedido=idpedido)
    return redirect('panel_inicio')


# ============================================================
# 6. PANEL - MENÚ
# ============================================================

def menu_panel(request):
    return render(request, 'panel/menu/inicioMenu.html')


# ------------------------------------------------------------
# Categorías
# ------------------------------------------------------------

def categoria_lista(request):
    return render(request, 'panel/menu/categorias/listaCategorias.html', {
        'categorias': CategoriaProducto.objects.all(),
    })


def categoria_crear(request):
    errores = []

    if request.method == 'POST':
        nombre = request.POST.get('nombre', '').strip()

        if not nombre:
            errores.append('El nombre es obligatorio.')
        else:
            CategoriaProducto.objects.create(nombre=nombre)
            messages.success(request, 'Categoría creada correctamente.')
            return redirect('categoria_lista')

    return render(request, 'panel/menu/categorias/crearCategoria.html', {
        'errores': errores,
    })


def categoria_editar(request, idcategoria):
    categoria = get_object_or_404(CategoriaProducto, pk=idcategoria)
    errores = []

    if request.method == 'POST':
        nombre = request.POST.get('nombre', '').strip()

        if not nombre:
            errores.append('El nombre es obligatorio.')
        else:
            categoria.nombre = nombre
            categoria.save()
            messages.success(request, 'Categoría actualizada correctamente.')
            return redirect('categoria_lista')

    return render(request, 'panel/menu/categorias/editarCategoria.html', {
        'categoria': categoria,
        'errores': errores,
    })


def categoria_eliminar(request, idcategoria):
    categoria = get_object_or_404(CategoriaProducto, pk=idcategoria)

    if request.method == 'POST':
        categoria.delete()
        messages.success(request, 'Categoría eliminada.')
        return redirect('categoria_lista')

    return render(request, 'panel/menu/categorias/eliminarCategoria.html', {
        'titulo': 'categoría',
        'objeto': categoria,
        'cancel_url': 'categoria_lista',
    })


# ------------------------------------------------------------
# Grupos de opciones
# ------------------------------------------------------------

def grupo_lista(request):
    return render(request, 'panel/menu/opciones/listaOpciones.html', {
        'grupos': GrupoOpcion.objects.prefetch_related('opcion_set').all(),
    })


def grupo_crear(request):
    errores = []
    insumos = Insumo.objects.all().order_by('nombreinsumo')

    if request.method == 'POST':
        nombre = request.POST.get('nombre', '').strip()
        minselecciones = request.POST.get('minselecciones') or 0
        maxselecciones = request.POST.get('maxselecciones') or 1
        opcion_nombres = request.POST.getlist('opcion_nombre[]')
        opcion_precios = request.POST.getlist('opcion_precio[]')
        opcion_insumos = request.POST.getlist('opcion_insumo[]')
        opcion_cantidades = request.POST.getlist('opcion_cantidad_insumo[]')

        if not nombre:
            errores.append('El nombre del grupo es obligatorio.')
        else:
            grupo = GrupoOpcion.objects.create(
                nombre=nombre,
                minselecciones=minselecciones,
                maxselecciones=maxselecciones,
            )

            total_opc = len(opcion_nombres)
            for i in range(total_opc):
                nom = opcion_nombres[i].strip() if i < len(opcion_nombres) else ''
                if not nom:
                    continue
                precio = opcion_precios[i] if i < len(opcion_precios) else 0
                insumo_id = opcion_insumos[i] if i < len(opcion_insumos) else None
                cant_ins = opcion_cantidades[i] if i < len(opcion_cantidades) else 1

                insumo_obj = Insumo.objects.filter(pk=insumo_id).first() if insumo_id and str(insumo_id).isdigit() else None
                try:
                    cant_dec = Decimal(str(cant_ins)) if cant_ins else Decimal('1.00')
                    if cant_dec <= 0:
                        cant_dec = Decimal('1.00')
                except Exception:
                    cant_dec = Decimal('1.00')

                Opcion.objects.create(
                    idgrupo=grupo,
                    nombre=nom,
                    precioadicional=precio or 0,
                    idinsumo=insumo_obj,
                    cantidad_insumo=cant_dec,
                )

            messages.success(
                request, 'Grupo de opciones creado correctamente.'
            )
            return redirect('grupo_lista')

    return render(request, 'panel/menu/opciones/crearGrupo.html', {
        'errores': errores,
        'insumos': insumos,
    })


def grupo_editar(request, idgrupo):
    grupo = get_object_or_404(GrupoOpcion, pk=idgrupo)
    errores = []
    insumos = Insumo.objects.all().order_by('nombreinsumo')

    if request.method == 'POST':
        nombre = request.POST.get('nombre', '').strip()

        grupo.minselecciones = request.POST.get('minselecciones') or 0
        grupo.maxselecciones = request.POST.get('maxselecciones') or 1

        if not nombre:
            errores.append('El nombre del grupo es obligatorio.')
        else:
            grupo.nombre = nombre
            grupo.save()

            opcion_ids = request.POST.getlist('opcion_id[]')
            opcion_nombres = request.POST.getlist('opcion_nombre[]')
            opcion_precios = request.POST.getlist('opcion_precio[]')
            opcion_insumos = request.POST.getlist('opcion_insumo[]')
            opcion_cantidades = request.POST.getlist('opcion_cantidad_insumo[]')
            opcion_eliminar = set(request.POST.getlist('opcion_eliminar[]'))

            if opcion_eliminar:
                Opcion.objects.filter(idgrupo=grupo, pk__in=opcion_eliminar).delete()

            total_opc = len(opcion_nombres)
            for i in range(total_opc):
                opc_id_raw = opcion_ids[i].strip() if i < len(opcion_ids) else ''
                if opc_id_raw in opcion_eliminar:
                    continue

                nom = opcion_nombres[i].strip() if i < len(opcion_nombres) else ''
                if not nom:
                    continue

                precio = opcion_precios[i] if i < len(opcion_precios) else 0
                insumo_id = opcion_insumos[i] if i < len(opcion_insumos) else None
                cant_ins = opcion_cantidades[i] if i < len(opcion_cantidades) else 1

                insumo_obj = Insumo.objects.filter(pk=insumo_id).first() if insumo_id and str(insumo_id).isdigit() else None
                try:
                    cant_dec = Decimal(str(cant_ins)) if cant_ins else Decimal('1.00')
                    if cant_dec <= 0:
                        cant_dec = Decimal('1.00')
                except Exception:
                    cant_dec = Decimal('1.00')

                if opc_id_raw and opc_id_raw.isdigit():
                    opcion_obj = Opcion.objects.filter(pk=int(opc_id_raw), idgrupo=grupo).first()
                    if opcion_obj:
                        opcion_obj.nombre = nom
                        opcion_obj.precioadicional = precio or 0
                        opcion_obj.idinsumo = insumo_obj
                        opcion_obj.cantidad_insumo = cant_dec
                        opcion_obj.save()
                else:
                    Opcion.objects.create(
                        idgrupo=grupo,
                        nombre=nom,
                        precioadicional=precio or 0,
                        idinsumo=insumo_obj,
                        cantidad_insumo=cant_dec,
                    )

            messages.success(request, 'Grupo actualizado correctamente.')
            return redirect('grupo_lista')

    return render(request, 'panel/menu/opciones/editarGrupo.html', {
        'grupo': grupo,
        'opciones': grupo.opcion_set.select_related('idinsumo').all(),
        'insumos': insumos,
        'errores': errores,
    })


def grupo_eliminar(request, idgrupo):
    grupo = get_object_or_404(GrupoOpcion, pk=idgrupo)

    if request.method == 'POST':
        grupo.delete()
        messages.success(request, 'Grupo eliminado correctamente.')
        return redirect('grupo_lista')

    return render(request, 'panel/menu/opciones/eliminarGrupo.html', {
        'titulo': 'grupo de opciones',
        'objeto': grupo,
        'cancel_url': 'grupo_lista',
    })


# ------------------------------------------------------------
# Productos
# ------------------------------------------------------------

def _ids_enviados(request, campo):
    return [int(v) for v in request.POST.getlist(campo) if v.isdigit()]


def _validar_producto(nombre, precio_raw):
    """Devuelve (precio, errores) para los datos básicos del producto."""
    errores = []

    if not nombre:
        errores.append('El nombre del producto es obligatorio.')

    precio = None

    try:
        precio = Decimal(precio_raw)

        if precio < Decimal('0.00'):
            errores.append('El precio no puede ser negativo.')

    except (InvalidOperation, ValueError, TypeError):
        errores.append('El precio debe ser un número válido.')

    return precio, errores


def _guardar_relaciones_producto(producto, grupos_ids, extras_ids):
    """Reemplaza los grupos de opciones y extras del producto."""
    ProductoGrupoOpcion.objects.filter(idproducto=producto).delete()

    for grupo in GrupoOpcion.objects.filter(pk__in=grupos_ids):
        ProductoGrupoOpcion.objects.create(idproducto=producto, idgrupo=grupo)

    ProductoExtras.objects.filter(idproducto=producto).delete()

    for extra in Extras.objects.filter(pk__in=extras_ids):
        ProductoExtras.objects.create(idproducto=producto, idextra=extra)


def producto_lista(request):
    productos = (
        Producto.objects
        .select_related('idcategoria', 'idestadoproducto')
        .prefetch_related(
            'productogrupoopcion_set__idgrupo',
            'productoextras_set__idextra',
        )
        .order_by('-idproducto')
    )

    return render(request, 'panel/menu/productos/listaProductos.html', {
        'productos': productos,
        'ESTADO_ELIMINADO': ESTADO_ELIMINADO,
        'ESTADO_OCULTO': ESTADO_OCULTO,
    })


def producto_crear(request):
    contexto = {
        'categorias': CategoriaProducto.objects.all(),
        'grupos': GrupoOpcion.objects.prefetch_related('opcion_set').all(),
        'extras': Extras.objects.all(),
        'errores': [],
        'datos_previos': {},
        'grupos_seleccionados': [],
        'extras_seleccionados': [],
    }

    if request.method != 'POST':
        return render(
            request, 'panel/menu/productos/crearProducto.html', contexto
        )

    nombre = (request.POST.get('nombre') or '').strip()
    precio_raw = (request.POST.get('precio') or '').strip()
    idcategoria = request.POST.get('idcategoria')
    grupos_ids = _ids_enviados(request, 'grupos')
    extras_ids = _ids_enviados(request, 'extras')

    precio, errores = _validar_producto(nombre, precio_raw)

    if errores:
        return render(request, 'panel/menu/productos/crearProducto.html', {
            **contexto,
            'errores': errores,
            'datos_previos': request.POST,
            'grupos_seleccionados': grupos_ids,
            'extras_seleccionados': extras_ids,
        })

    categoria = (
        CategoriaProducto.objects.filter(pk=idcategoria).first()
        if idcategoria else None
    )

    with transaction.atomic():
        producto = Producto.objects.create(
            nombre=nombre,
            descripcion=request.POST.get('descripcion', ''),
            precio=precio,
            idcategoria=categoria,
            idnegocio=Negocio.objects.first(),
            imagen=request.FILES.get('imagen'),
        )

        _guardar_relaciones_producto(producto, grupos_ids, extras_ids)

    messages.success(
        request, f'Producto "{producto.nombre}" creado exitosamente.'
    )
    return redirect('producto_lista')


def producto_editar(request, idproducto):
    producto = get_object_or_404(Producto, pk=idproducto)

    contexto = {
        'producto': producto,
        'categorias': CategoriaProducto.objects.all(),
        'grupos': GrupoOpcion.objects.prefetch_related('opcion_set').all(),
        'extras': Extras.objects.all(),
    }

    if request.method != 'POST':
        return render(request, 'panel/menu/productos/editarProducto.html', {
            **contexto,
            'errores': [],
            'grupos_seleccionados': list(
                ProductoGrupoOpcion.objects
                .filter(idproducto=producto)
                .values_list('idgrupo_id', flat=True)
            ),
            'extras_seleccionados': list(
                ProductoExtras.objects
                .filter(idproducto=producto)
                .values_list('idextra_id', flat=True)
            ),
        })

    nombre = (request.POST.get('nombre') or '').strip()
    descripcion = request.POST.get('descripcion', '').strip()
    precio_raw = (request.POST.get('precio') or '').strip()
    idcategoria = request.POST.get('idcategoria')
    grupos_ids = _ids_enviados(request, 'grupos')
    extras_ids = _ids_enviados(request, 'extras')

    precio, errores = _validar_producto(nombre, precio_raw)

    if errores:
        return render(request, 'panel/menu/productos/editarProducto.html', {
            **contexto,
            'errores': errores,
            'grupos_seleccionados': grupos_ids,
            'extras_seleccionados': extras_ids,
        })

    with transaction.atomic():
        producto.nombre = nombre
        producto.descripcion = descripcion
        producto.precio = precio
        producto.idcategoria = (
            CategoriaProducto.objects.filter(pk=idcategoria).first()
            if idcategoria else None
        )

        # Si no se sube una imagen nueva se conserva la actual.
        if request.FILES.get('imagen'):
            producto.imagen = request.FILES['imagen']

        producto.save()
        _guardar_relaciones_producto(producto, grupos_ids, extras_ids)

    messages.success(
        request, f'Producto "{producto.nombre}" actualizado correctamente.'
    )
    return redirect('producto_lista')


def producto_eliminar(request, idproducto):
    """
    Borrado lógico: el producto cambia a estado Eliminado (o vuelve a Activo
    si ya estaba eliminado). No se borra el registro porque lo referencian
    recetas y pedidos históricos.
    """
    producto = get_object_or_404(Producto, pk=idproducto)
    esta_eliminado = producto.idestadoproducto_id == ESTADO_ELIMINADO

    if request.method == 'POST':
        nuevo_estado = ESTADO_ACTIVO if esta_eliminado else ESTADO_ELIMINADO

        # update() evita ejecutar el save() del modelo (conversión de imagen).
        Producto.objects.filter(pk=producto.pk).update(
            idestadoproducto=nuevo_estado
        )

        messages.success(
            request,
            f'Producto "{producto.nombre}" '
            f'{"restaurado" if esta_eliminado else "eliminado"}.'
        )
        return redirect('producto_lista')

    return render(request, 'panel/menu/productos/eliminarProducto.html', {
        'producto': producto,
        'esta_eliminado': esta_eliminado,
        'cancel_url': 'producto_lista',
    })


# ------------------------------------------------------------
# Extras
# ------------------------------------------------------------

def extra_lista(request):
    extras_con_conteo = [
        {
            'extra': extra,
            'cant_productos': ProductoExtras.objects.filter(
                idextra=extra
            ).count(),
        }
        for extra in Extras.objects.all().order_by('idextra')
    ]

    return render(request, 'panel/menu/extras/listaExtras.html', {
        'extras_con_conteo': extras_con_conteo,
    })


def extra_crear(request):
    if request.method == 'POST':
        form = ExtraForm(request.POST)

        if form.is_valid():
            extra = form.save()
            messages.success(
                request, f'Extra "{extra.nombre}" creado exitosamente.'
            )
            return redirect('extra_lista')
    else:
        form = ExtraForm()

    return render(request, 'panel/menu/extras/crearExtra.html', {
        'form': form,
        'errores': [],
    })


def extra_editar(request, idextra):
    extra = get_object_or_404(Extras, pk=idextra)

    if request.method == 'POST':
        form = ExtraForm(request.POST, instance=extra)

        if form.is_valid():
            extra = form.save()
            messages.success(
                request, f'Extra "{extra.nombre}" actualizado correctamente.'
            )
            return redirect('extra_lista')
    else:
        form = ExtraForm(instance=extra)

    return render(request, 'panel/menu/extras/editarExtra.html', {
        'extra': extra,
        'form': form,
        'errores': [],
    })


def extra_eliminar(request, idextra):
    extra = get_object_or_404(Extras, pk=idextra)
    cant_productos = ProductoExtras.objects.filter(idextra=extra).count()

    if request.method == 'POST':
        nombre_extra = extra.nombre

        with transaction.atomic():
            ProductoExtras.objects.filter(idextra=extra).delete()
            extra.delete()

        messages.success(request, f'Extra "{nombre_extra}" eliminado.')
        return redirect('extra_lista')

    return render(request, 'panel/menu/extras/eliminarExtra.html', {
        'extra': extra,
        'cant_productos': cant_productos,
        'cancel_url': 'extra_lista',
    })


# ------------------------------------------------------------
# Recetas
# ------------------------------------------------------------

def gestionar_receta(request, idproducto):
    producto = get_object_or_404(Producto, pk=idproducto)
    receta, _ = Receta.objects.get_or_create(idproducto=producto)

    origen = request.GET.get('origen') or request.POST.get('origen', '')

    if request.method == 'POST':
        id_insumo = request.POST.get('id_insumo')
        cantidad_texto = request.POST.get('cantidadinsumo', '').strip()
        es_removible = request.POST.get('es_removible') == 'on'

        if not id_insumo:
            messages.error(request, 'Tenés que seleccionar un insumo.')

        else:
            try:
                cantidad = Decimal(cantidad_texto)

                if (
                    not cantidad.is_finite()
                    or cantidad <= 0
                    or cantidad.as_tuple().exponent < -2
                ):
                    raise InvalidOperation

            except (InvalidOperation, ValueError):
                messages.error(
                    request,
                    'Ingresá una cantidad mayor que cero, '
                    'con hasta dos decimales.'
                )

            else:
                insumo = get_object_or_404(Insumo, pk=id_insumo)

                existe = RecetaInsumo.objects.filter(
                    idreceta=receta,
                    idinsumo=insumo,
                ).exists()

                if existe:
                    messages.warning(
                        request,
                        f'El insumo "{insumo.nombreinsumo}" '
                        'ya está en esta receta.'
                    )

                else:
                    try:
                        with transaction.atomic():
                            RecetaInsumo.objects.create(
                                idreceta=receta,
                                idinsumo=insumo,
                                cantidadinsumo=cantidad,
                                unidadmedida=insumo.unidadmedidaegreso,
                                es_removible=es_removible,
                            )

                        messages.success(
                            request,
                            f'"{insumo.nombreinsumo}" '
                            'se agregó correctamente a la receta.'
                        )

                    except IntegrityError:
                        messages.error(
                            request,
                            'No se pudo guardar el ingrediente. '
                            'Verificá que la tabla RecetaInsumo '
                            'tenga las columnas correctas en SQL Server.'
                        )

        url = reverse(
            'gestionar_receta',
            kwargs={'idproducto': producto.idproducto},
        )

        if origen:
            url += f'?origen={quote(origen)}'

        return redirect(url)

    detalles_receta = (
        RecetaInsumo.objects
        .filter(idreceta=receta)
        .select_related('idinsumo')
    )

    return render(request, 'panel/menu/productos/gestionarReceta.html', {
        'producto': producto,
        'receta': receta,
        'detalles_receta': detalles_receta,
        'todos_los_insumos': Insumo.objects.all(),
        'origen': origen,
    })


def eliminar_insumo_receta(request, idrecetainsumo):
    detalle = get_object_or_404(RecetaInsumo, pk=idrecetainsumo)

    id_producto = detalle.idreceta.idproducto.idproducto
    origen = request.GET.get('origen', '')

    nombre_insumo = (
        detalle.idinsumo.nombreinsumo if detalle.idinsumo else 'Insumo'
    )

    if request.method == 'POST':
        detalle.delete()
        messages.success(
            request,
            f'Insumo "{nombre_insumo}" quitado de la receta.'
        )
    else:
        messages.warning(
            request,
            'Para quitar un ingrediente, se necesita una solicitud POST.'
        )

    url = reverse('gestionar_receta', kwargs={'idproducto': id_producto})

    if origen:
        url += f'?origen={quote(origen)}'

    return redirect(url)


def editar_cantidad_insumo_receta(request, idrecetainsumo):
    detalle = get_object_or_404(
        RecetaInsumo.objects.select_related(
            'idreceta',
            'idreceta__idproducto',
            'idinsumo',
        ),
        pk=idrecetainsumo,
    )

    producto = detalle.idreceta.idproducto

    origen = request.POST.get('origen', '') or request.GET.get('origen', '')

    if request.method != 'POST':
        return redirect('gestionar_receta', idproducto=producto.idproducto)

    cantidad_texto = request.POST.get('cantidadinsumo', '').strip()

    try:
        cantidad = Decimal(cantidad_texto)

        if (
            not cantidad.is_finite()
            or cantidad <= 0
            or cantidad.as_tuple().exponent < -2
        ):
            raise InvalidOperation

    except (InvalidOperation, ValueError):
        messages.error(
            request,
            'Ingresá una cantidad mayor que cero, con hasta dos decimales.'
        )

    else:
        detalle.cantidadinsumo = cantidad
        detalle.save(update_fields=['cantidadinsumo'])

        nombre_insumo = (
            detalle.idinsumo.nombreinsumo if detalle.idinsumo
            else 'Ingrediente'
        )

        messages.success(
            request,
            f'La cantidad de "{nombre_insumo}" se actualizó correctamente.'
        )

    url = reverse(
        'gestionar_receta',
        kwargs={'idproducto': producto.idproducto},
    )

    if origen:
        url += f'?origen={quote(origen)}'

    return redirect(url)


# ============================================================
# 7. PANEL - PROMOCIONES
# ============================================================

TEMPLATE_FORM_PROMOCION = 'panel/promociones/formPromocion.html'


def _valor_beneficio(tipo_beneficio, valor_raw):
    """Los beneficios sin valor numérico (ej. 2x1) se guardan como NULL."""
    if tipo_beneficio.comportamiento == 'SIN_VALOR' or not valor_raw:
        return None

    return valor_raw


def _form_promocion_con_error(request, contexto, mensaje):
    return render(request, TEMPLATE_FORM_PROMOCION, {
        **contexto,
        'error': mensaje,
    })


def promociones_panel(request):
    return render(request, 'panel/promociones/inicioPromociones.html')


def lista_promociones(request):
    return render(request, 'panel/promociones/listaPromociones.html', {
        'promociones': Promocion.objects.all(),
    })


def crear_promocion(request):
    contexto = {
        'titulo': 'Crear promoción',
        'tipos_beneficio': TipoBeneficio.objects.all(),
        'productos': Producto.objects.all(),
    }

    if request.method != 'POST':
        return render(request, TEMPLATE_FORM_PROMOCION, contexto)

    palabraclave = request.POST.get('palabraclave', '').strip()
    tipobeneficio_id = request.POST.get('tipobeneficio')
    producto_id = request.POST.get('producto')
    descuento = request.POST.get('descuento', '').strip()
    fechainicio = request.POST.get('fechainicio') or None
    fechafin = request.POST.get('fechafin') or None
    activo = 1 if request.POST.get('activo') else 0

    if not palabraclave:
        return _form_promocion_con_error(
            request, contexto, 'Ingresá una palabra clave.'
        )

    if not tipobeneficio_id:
        return _form_promocion_con_error(
            request, contexto, 'Seleccioná un tipo de beneficio.'
        )

    if not producto_id:
        return _form_promocion_con_error(
            request, contexto, 'Seleccioná un producto.'
        )

    try:
        tipo_beneficio = TipoBeneficio.objects.get(pk=tipobeneficio_id)
    except TipoBeneficio.DoesNotExist:
        return _form_promocion_con_error(
            request, contexto, 'El tipo de beneficio seleccionado no existe.'
        )

    try:
        producto = Producto.objects.get(pk=producto_id)
    except Producto.DoesNotExist:
        return _form_promocion_con_error(
            request, contexto, 'El producto seleccionado no existe.'
        )

    valor = None

    if getattr(tipo_beneficio, 'comportamiento', None) != 'SIN_VALOR':
        if not descuento:
            return _form_promocion_con_error(
                request, contexto, 'Ingresá el valor del beneficio.'
            )

        try:
            valor = Decimal(descuento)
        except (InvalidOperation, TypeError):
            return _form_promocion_con_error(
                request, contexto, 'El valor del beneficio no es válido.'
            )

        if valor < 0:
            return _form_promocion_con_error(
                request, contexto,
                'El valor del beneficio no puede ser negativo.'
            )

    promocion = Promocion(
        palabraclave=palabraclave,
        tipobeneficio=tipo_beneficio,
        valor=valor,
        fechainicio=fechainicio,
        fechafin=fechafin,
        activo=activo,
    )
    promocion.save()

    ProductoPromocion.objects.create(
        idproducto=producto,
        idpromocion=promocion,
    )

    messages.success(
        request,
        f'La promoción "{palabraclave}" fue creada correctamente.'
    )
    return redirect('panel_lista_promociones')


def editar_promocion(request, idpromocion):
    promocion = get_object_or_404(Promocion, idpromocion=idpromocion)

    relacion_producto = (
        ProductoPromocion.objects
        .filter(idpromocion=promocion)
        .select_related('idproducto')
        .first()
    )

    contexto = {
        'titulo': 'Editar promoción',
        'promocion': promocion,
        'tipos_beneficio': TipoBeneficio.objects.all(),
        'productos': Producto.objects.all(),
        'tipo_beneficio_seleccionado': (
            promocion.tipobeneficio.pk if promocion.tipobeneficio else None
        ),
        'producto_seleccionado': (
            relacion_producto.idproducto.pk
            if relacion_producto and relacion_producto.idproducto
            else None
        ),
    }

    if request.method != 'POST':
        return render(request, TEMPLATE_FORM_PROMOCION, contexto)

    palabraclave = request.POST.get('palabraclave', '').strip()
    tipobeneficio_id = request.POST.get('tipobeneficio')
    producto_id = request.POST.get('producto')
    descuento = request.POST.get('descuento', '').strip()
    fechainicio = request.POST.get('fechainicio') or None
    fechafin = request.POST.get('fechafin') or None
    activo = 1 if request.POST.get('activo') else 0

    if not palabraclave:
        return _form_promocion_con_error(
            request, contexto, 'Ingresá una palabra clave.'
        )

    if not tipobeneficio_id:
        return _form_promocion_con_error(
            request, contexto, 'Seleccioná un tipo de beneficio.'
        )

    if not producto_id:
        return _form_promocion_con_error(
            request, contexto, 'Seleccioná un producto.'
        )

    try:
        tipo_beneficio = TipoBeneficio.objects.get(pk=tipobeneficio_id)
    except TipoBeneficio.DoesNotExist:
        return _form_promocion_con_error(
            request, contexto, 'El tipo de beneficio seleccionado no existe.'
        )

    try:
        producto = Producto.objects.get(pk=producto_id)
    except Producto.DoesNotExist:
        return _form_promocion_con_error(
            request, contexto, 'El producto seleccionado no existe.'
        )

    valor = None

    if getattr(tipo_beneficio, 'comportamiento', None) != 'SIN_VALOR':
        if not descuento:
            return _form_promocion_con_error(
                request, contexto, 'Ingresá el valor del beneficio.'
            )

        try:
            valor = Decimal(descuento)
        except (InvalidOperation, TypeError):
            return _form_promocion_con_error(
                request, contexto, 'El valor del beneficio no es válido.'
            )

        if valor < 0:
            return _form_promocion_con_error(
                request, contexto,
                'El valor del beneficio no puede ser negativo.'
            )

    promocion.palabraclave = palabraclave
    promocion.tipobeneficio = tipo_beneficio
    promocion.valor = valor
    promocion.fechainicio = fechainicio
    promocion.fechafin = fechafin
    promocion.activo = activo
    promocion.save()

    ProductoPromocion.objects.filter(idpromocion=promocion).delete()
    ProductoPromocion.objects.create(
        idproducto=producto,
        idpromocion=promocion,
    )

    messages.success(
        request,
        f'La promoción "{palabraclave}" fue actualizada correctamente.'
    )
    return redirect('panel_lista_promociones')


def eliminar_promocion(request, idpromocion):
    promocion = get_object_or_404(Promocion, pk=idpromocion)

    if request.method == 'POST':
        with transaction.atomic():
            ProductoPromocion.objects.filter(idpromocion=promocion).delete()
            promocion.delete()

        return redirect('lista_promociones')

    return render(request, 'panel/promociones/EliminarPromocion.html', {
        'promocion': promocion,
    })


# ------------------------------------------------------------
# Tipos de beneficio
# ------------------------------------------------------------

def lista_tipos_beneficio(request):
    return render(
        request,
        'panel/promociones/tiposBeneficio/listaTiposBeneficios.html',
        {'tipos': TipoBeneficio.objects.all()},
    )


def crear_tipo_beneficio(request):
    if request.method == 'GET':
        return render(
            request,
            'panel/promociones/tiposBeneficio/crearTipoBeneficio.html',
        )

    codigo = request.POST.get('codigo', '').strip()
    nombre = request.POST.get('nombre', '').strip()
    descripcion = request.POST.get('descripcion', '').strip()
    comportamiento = request.POST.get('comportamiento', '').strip()

    if not nombre:
        messages.error(request, 'Debes ingresar un nombre para el beneficio.')
        return redirect('crear_tipo_beneficio')

    if not comportamiento:
        messages.error(request, 'Debes seleccionar un comportamiento.')
        return redirect('crear_tipo_beneficio')

    # Si no se indica código, se genera a partir del nombre.
    if not codigo:
        codigo = nombre.upper().replace(' ', '_')

    if TipoBeneficio.objects.filter(codigo__iexact=codigo).exists():
        messages.error(
            request, 'Ya existe un tipo de beneficio con ese código.'
        )
        return redirect('crear_tipo_beneficio')

    if TipoBeneficio.objects.filter(nombre__iexact=nombre).exists():
        messages.error(
            request, 'Ya existe un tipo de beneficio con ese nombre.'
        )
        return redirect('crear_tipo_beneficio')

    TipoBeneficio.objects.create(
        codigo=codigo,
        nombre=nombre,
        descripcion=descripcion,
        comportamiento=comportamiento,
    )

    messages.success(request, 'Tipo de beneficio creado correctamente.')
    return redirect('lista_tipos_beneficio')


def editar_tipo_beneficio(request, pk):
    tipo = get_object_or_404(TipoBeneficio, pk=pk)

    if request.method == 'POST':
        codigo = request.POST.get('codigo', '').strip()
        nombre = request.POST.get('nombre', '').strip()
        descripcion = request.POST.get('descripcion', '').strip()

        if not nombre:
            messages.error(request, 'El nombre no puede estar vacío.')
            return redirect('editar_tipo_beneficio', pk=pk)

        if not codigo:
            codigo = nombre.upper().replace(' ', '_')

        # Se excluye el propio registro para poder conservar sus valores.
        if (
            TipoBeneficio.objects
            .filter(codigo__iexact=codigo)
            .exclude(pk=pk)
            .exists()
        ):
            messages.error(request, 'Ya existe otro beneficio con ese código.')
            return redirect('editar_tipo_beneficio', pk=pk)

        if (
            TipoBeneficio.objects
            .filter(nombre__iexact=nombre)
            .exclude(pk=pk)
            .exists()
        ):
            messages.error(request, 'Ya existe otro beneficio con ese nombre.')
            return redirect('editar_tipo_beneficio', pk=pk)

        tipo.codigo = codigo
        tipo.nombre = nombre
        tipo.descripcion = descripcion
        tipo.save()

        messages.success(
            request, 'Tipo de beneficio modificado correctamente.'
        )
        return redirect('lista_tipos_beneficio')

    return render(
        request,
        'panel/promociones/tiposBeneficio/editarTipoBeneficio.html',
        {'tipo': tipo},
    )


def eliminar_tipo_beneficio(request, pk):
    tipo = get_object_or_404(TipoBeneficio, pk=pk)

    if request.method == 'GET':
        return render(
            request,
            'panel/promociones/tiposBeneficio/eliminarTipoBeneficio.html',
            {'tipo': tipo},
        )

    if Promocion.objects.filter(tipobeneficio=tipo).exists():
        messages.error(
            request,
            'No puedes eliminar este beneficio porque tiene '
            'promociones asociadas.'
        )
        return redirect('panel_lista_tipos_beneficio')

    tipo.delete()

    messages.success(request, 'Tipo de beneficio eliminado correctamente.')
    return redirect('panel_lista_tipos_beneficio')


# ============================================================
# 8. PANEL - STOCK
# ============================================================


def stock_lista(request):
    hoy = timezone.localdate()
    fecha_limite = hoy + timedelta(days=7)

    insumos = Insumo.objects.all().order_by('nombreinsumo')

    # Stock en el mínimo o por debajo.
    insumos_criticos = [
        insumo for insumo in insumos
        if insumo.stockactual <= insumo.stockminimo
    ]

    # Lotes que vencen entre hoy y los próximos 7 días.
    lotes_por_vencer = (
        Lote.objects
        .filter(
            fechavencimiento__gte=hoy,
            fechavencimiento__lte=fecha_limite,
            cantidadactual__gt=0,
        )
        .select_related('idinsumo')
        .order_by('fechavencimiento', 'idinsumo__nombreinsumo')
    )

    # Lotes que ya vencieron y todavía tienen cantidad registrada.
    lotes_vencidos = (
        Lote.objects
        .filter(
            fechavencimiento__lt=hoy,
            cantidadactual__gt=0,
        )
        .select_related('idinsumo')
        .order_by('fechavencimiento', 'idinsumo__nombreinsumo')
    )

    # Se conserva el listado completo de lotes.
    lotes = (
        Lote.objects
        .select_related('idinsumo')
        .order_by('-fechaingreso', '-idlote')
    )

    compras = (
        Compras.objects
        .prefetch_related('detallecompra_set__idinsumo')
        .order_by('-fecha', '-idcompra')
    )

    return render(request, 'panel/stock/stockLista.html', {
        'insumos': insumos,
        'insumos_criticos': insumos_criticos,
        'lotes': lotes,
        'compras': compras,
        'lotes_por_vencer': lotes_por_vencer,
        'lotes_vencidos': lotes_vencidos,
        'hoy': hoy,
        'fecha_limite': fecha_limite,
    })



def stock_crear(request):
    if request.method == 'POST':
        Insumo.objects.create(
            nombreinsumo=request.POST.get('nombreinsumo'),
            unidadmedidaingreso=request.POST.get('unidadmedidaingreso'),
            unidadmedidaegreso=request.POST.get('unidadmedidaegreso'),
            stockactual=request.POST.get('stockactual', 0),
            stockminimo=request.POST.get('stockminimo', 5),
        )

        messages.success(
            request, 'Insumo creado correctamente en el stock.'
        )
        return redirect('stock_lista')

    return render(request, 'panel/stock/stockForm.html')


def stock_editar(request, idinsumo):
    insumo = get_object_or_404(Insumo, pk=idinsumo)

    if request.method == 'POST':
        insumo.nombreinsumo = request.POST.get('nombreinsumo', '').strip()
        insumo.unidadmedidaingreso = request.POST.get(
            'unidadmedidaingreso', ''
        ).strip()
        insumo.unidadmedidaegreso = request.POST.get(
            'unidadmedidaegreso', ''
        ).strip()
        insumo.stockactual = request.POST.get('stockactual') or 0
        insumo.stockminimo = request.POST.get('stockminimo') or 5
        insumo.descontar_stock_automaticamente = (
            request.POST.get('descontar_stock_automaticamente') == 'on'
        )

        if (
            not insumo.nombreinsumo
            or not insumo.unidadmedidaingreso
            or not insumo.unidadmedidaegreso
        ):
            messages.error(
                request,
                'Completá el nombre y las unidades de medida del insumo.'
            )
            return render(
                request, 'panel/stock/editarStock.html', {'insumo': insumo}
            )

        insumo.save()

        messages.success(request, 'Insumo actualizado correctamente.')
        return redirect('stock_lista')

    return render(request, 'panel/stock/editarStock.html', {'insumo': insumo})


def registrar_compra_insumo(request):
    if request.method == 'POST':
        insumo = get_object_or_404(Insumo, pk=request.POST.get('idinsumo'))

        try:
            cantidad = Decimal(request.POST.get('cantidad', '0'))
            costo_unitario = Decimal(request.POST.get('costounitario', '0'))
        except (ValueError, TypeError, ArithmeticError):
            messages.error(request, 'Ingresa cantidades y costos válidos.')
            return redirect('registrar_compra_insumo')

        if (
            not cantidad.is_finite()
            or not costo_unitario.is_finite()
            or cantidad <= 0
            or costo_unitario < 0
        ):
            messages.error(
                request,
                'La cantidad debe ser mayor que cero y el costo no puede '
                'ser negativo.',
            )
            return redirect('registrar_compra_insumo')

        fecha_compra = datetime.now()
        precio_total = cantidad * costo_unitario

        with transaction.atomic():
            insumo = Insumo.objects.select_for_update().get(pk=insumo.pk)

            compra = Compras.objects.create(
                fecha=fecha_compra,
                preciototal=precio_total,
            )

            DetalleCompra.objects.create(
                idcompra=compra,
                idinsumo=insumo,
                cantidad=cantidad,
                costounitario=costo_unitario,
            )

            Lote.objects.create(
                idinsumo=insumo,
                fechaingreso=fecha_compra,
                cantidadinicial=cantidad,
                cantidadactual=cantidad,
            )

            insumo.stockactual += cantidad
            insumo.save(update_fields=['stockactual'])

        messages.success(
            request,
            'Compra registrada, lote creado y stock actualizado '
            'correctamente.',
        )
        return redirect('stock_lista')

    return render(request, 'panel/stock/registrarCompra.html', {
        'insumos': Insumo.objects.all(),
    })


def ajustar_stock(request, idinsumo):
    insumo = get_object_or_404(Insumo, pk=idinsumo)

    if request.method == 'POST':
        cantidad_nueva_texto = request.POST.get('cantidadnueva', '').strip()
        motivo = request.POST.get('motivo', '').strip()

        try:
            cantidad_nueva = Decimal(cantidad_nueva_texto)

            if not cantidad_nueva.is_finite() or cantidad_nueva < 0:
                raise InvalidOperation

        except (InvalidOperation, ValueError):
            messages.error(
                request,
                'Ingresa una cantidad válida, igual o mayor que cero.'
            )
            return redirect('ajustar_stock', idinsumo=idinsumo)

        if not motivo:
            messages.error(request, 'Debes indicar el motivo del ajuste.')
            return redirect('ajustar_stock', idinsumo=idinsumo)

        if len(motivo) > 250:
            messages.error(
                request,
                'El motivo no puede superar los 250 caracteres.'
            )
            return redirect('ajustar_stock', idinsumo=idinsumo)

        with transaction.atomic():
            insumo = Insumo.objects.select_for_update().get(pk=idinsumo)
            cantidad_anterior = insumo.stockactual

            AjusteStock.objects.create(
                idinsumo=insumo,
                cantidadanterior=cantidad_anterior,
                cantidadnueva=cantidad_nueva,
                motivo=motivo,
            )

            insumo.stockactual = cantidad_nueva
            insumo.save(update_fields=['stockactual'])

        messages.success(request, 'Stock ajustado correctamente.')
        return redirect('stock_lista')

    return render(request, 'panel/stock/ajustarStock.html', {
        'insumo': insumo,
    })


# ============================================================
# 9. PANEL - USUARIOS
# ============================================================

def _error_correo(correo, idempleado=None):
    """Devuelve un mensaje de error o None si el correo es válido."""
    if not correo:
        return 'Debes indicar el correo electrónico.'

    if len(correo) > 254:
        return 'El correo no puede superar los 254 caracteres.'

    try:
        validate_email(correo)
    except ValidationError:
        return 'Ingresa un correo electrónico válido.'

    repetidos = Empleado.objects.filter(correo__iexact=correo)

    if idempleado is not None:
        repetidos = repetidos.exclude(pk=idempleado)

    if repetidos.exists():
        return 'Ya existe un usuario con ese correo.'

    return None


def empleado_lista(request):
    if request.method == 'POST':
        nombre = request.POST.get('nombre', '').strip()
        apellido = request.POST.get('apellido', '').strip()
        idrol = request.POST.get('idrol', '').strip()
        correo = request.POST.get('correo', '').strip().lower()
        contrasena = request.POST.get('contrasena', '')
        repetida = request.POST.get('confirmar_contrasena', '')

        if not nombre or not apellido:
            messages.error(request, 'Debes indicar el nombre y el apellido.')
            return redirect('empleado_lista')

        if len(nombre) > 100 or len(apellido) > 100:
            messages.error(
                request,
                'El nombre y el apellido no pueden superar los 100 caracteres.'
            )
            return redirect('empleado_lista')

        error_correo = _error_correo(correo)
        if error_correo:
            messages.error(request, error_correo)
            return redirect('empleado_lista')

        try:
            rol = Rol.objects.get(pk=idrol)
        except (Rol.DoesNotExist, ValueError, TypeError):
            messages.error(request, 'Selecciona un rol válido.')
            return redirect('empleado_lista')

        if not contrasena:
            messages.error(request, 'Debes indicar una contraseña.')
            return redirect('empleado_lista')

        if contrasena != repetida:
            messages.error(request, 'Las contraseñas no coinciden.')
            return redirect('empleado_lista')

        if len(contrasena) < 8:
            messages.error(
                request,
                'La contraseña debe tener al menos 8 caracteres.'
            )
            return redirect('empleado_lista')

        with transaction.atomic():
            Empleado.objects.create(
                nombre=nombre,
                apellido=apellido,
                correo=correo,
                id_rol=rol.pk,
                contrasena=make_password(contrasena),
            )

        messages.success(request, 'Usuario creado correctamente.')
        return redirect('empleado_lista')

    empleados = Empleado.objects.order_by('apellido', 'nombre')

    return render(
        request,
        'panel/administracion/usuarios/empleadoLista.html',
        {
            'empleados': empleados,
            'roles': Rol.objects.all(),
        },
    )


def empleado_crear(request):
    if request.method == 'POST':
        nombre = request.POST.get('nombre', '').strip()
        apellido = request.POST.get('apellido', '').strip()
        correo = request.POST.get('correo', '').strip().lower()
        idrol = request.POST.get('idrol', '').strip()
        contrasena = request.POST.get('contrasena', '')
        repetida = request.POST.get('confirmar_contrasena', '')

        if not nombre or not apellido:
            messages.error(request, 'Debes indicar el nombre y el apellido.')
            return redirect('empleado_crear')

        if len(nombre) > 100 or len(apellido) > 100:
            messages.error(
                request,
                'El nombre y el apellido no pueden superar los 100 caracteres.'
            )
            return redirect('empleado_crear')

        error_correo = _error_correo(correo)
        if error_correo:
            messages.error(request, error_correo)
            return redirect('empleado_crear')

        try:
            rol = Rol.objects.get(pk=idrol)
        except (Rol.DoesNotExist, ValueError, TypeError):
            messages.error(request, 'Seleccioná un rol válido.')
            return redirect('empleado_crear')

        if not contrasena:
            messages.error(request, 'Debes indicar una contraseña.')
            return redirect('empleado_crear')

        if contrasena != repetida:
            messages.error(request, 'Las contraseñas no coinciden.')
            return redirect('empleado_crear')

        if len(contrasena) < 8:
            messages.error(
                request,
                'La contraseña debe tener al menos 8 caracteres.'
            )
            return redirect('empleado_crear')

        with transaction.atomic():
            Empleado.objects.create(
                nombre=nombre,
                apellido=apellido,
                correo=correo,
                id_rol=rol.pk,
                contrasena=make_password(contrasena),
            )

        messages.success(request, 'Usuario creado correctamente.')
        return redirect('empleado_lista')

    return render(
        request,
        'panel/administracion/usuarios/empleadoCrear.html',
        {'roles': Rol.objects.all()},
    )

def empleado_editar(request, idempleado):
    empleado = get_object_or_404(Empleado, pk=idempleado)

    if request.method == 'POST':
        nombre = request.POST.get('nombre', '').strip()
        apellido = request.POST.get('apellido', '').strip()
        idrol = request.POST.get('idrol', '').strip()
        correo = request.POST.get('correo', '').strip().lower()
        contrasena = request.POST.get('contrasena', '')
        repetida = request.POST.get('confirmar_contrasena', '')

        if not nombre or not apellido:
            messages.error(
                request,
                'Debes indicar el nombre y el apellido.'
            )
            return redirect(
                'empleado_editar',
                idempleado=empleado.pk
            )

        if len(nombre) > 100 or len(apellido) > 100:
            messages.error(
                request,
                'El nombre y el apellido no pueden superar los 100 caracteres.'
            )
            return redirect(
                'empleado_editar',
                idempleado=empleado.pk
            )

        error_correo = _error_correo(correo, empleado.pk)

        if error_correo:
            messages.error(request, error_correo)
            return redirect(
                'empleado_editar',
                idempleado=empleado.pk
            )

        try:
            rol = Rol.objects.get(pk=idrol)
        except (Rol.DoesNotExist, ValueError, TypeError):
            messages.error(request, 'Seleccioná un rol válido.')
            return redirect(
                'empleado_editar',
                idempleado=empleado.pk
            )

        if contrasena or repetida:
            if contrasena != repetida:
                messages.error(
                    request,
                    'Las contraseñas no coinciden.'
                )
                return redirect(
                    'empleado_editar',
                    idempleado=empleado.pk
                )

            if len(contrasena) < 8:
                messages.error(
                    request,
                    'La contraseña debe tener al menos 8 caracteres.'
                )
                return redirect(
                    'empleado_editar',
                    idempleado=empleado.pk
                )

        with transaction.atomic():
            empleado = Empleado.objects.select_for_update().get(
                pk=idempleado
            )

            empleado.nombre = nombre
            empleado.apellido = apellido
            empleado.correo = correo
            empleado.id_rol = rol.pk

            campos = [
                'nombre',
                'apellido',
                'correo',
                'id_rol',
            ]

            if contrasena:
                empleado.contrasena = make_password(contrasena)
                campos.append('contrasena')

            empleado.save(update_fields=campos)

        messages.success(
            request,
            'Usuario actualizado correctamente.'
        )
        return redirect('empleado_lista')

    return render(
        request,
        'panel/administracion/usuarios/empleadoModificar.html',
        {
            'empleado': empleado,
            'roles': Rol.objects.all(),
        },
    )

def empleado_eliminar(request, idempleado):
    empleado = get_object_or_404(Empleado, pk=idempleado)

    # Si se confirma el formulario, se elimina el usuario.
    if request.method == 'POST':
        nombre_completo = f'{empleado.nombre} {empleado.apellido}'

        try:
            empleado.delete()

            messages.success(
                request,
                f'El usuario "{nombre_completo}" fue eliminado correctamente.'
            )

        except IntegrityError:
            messages.error(
                request,
                'No se puede eliminar este usuario porque tiene registros relacionados.'
            )

        return redirect('empleado_lista')

    # Si se ingresa mediante GET, muestra la página de confirmación.
    return render(
        request,
        'panel/administracion/usuarios/empleadoEliminar.html',
        {'empleado': empleado}
    )


# ============================================================
# 10. PANEL - ROLES
# ============================================================

def rol_lista(request):
    roles = Rol.objects.all().order_by('nombrerol')

    return render(
        request,
        'panel/administracion/roles/rolLista.html',
        {'roles': roles},
    )


def rol_crear(request):
    permisos = Permiso.objects.all().order_by('modulo', 'accion')
    permisos_seleccionados = []

    if request.method == 'POST':
        nombre = request.POST.get('nombrerol', '').strip()

        for valor in request.POST.getlist('permisos'):
            try:
                permisos_seleccionados.append(int(valor))
            except (ValueError, TypeError):
                continue

        if not nombre:
            messages.error(request, 'El nombre del rol es obligatorio.')

        elif len(nombre) > 50:
            messages.error(
                request,
                'El nombre del rol no puede superar los 50 caracteres.'
            )

        elif Rol.objects.filter(nombrerol__iexact=nombre).exists():
            messages.error(request, 'Ya existe un rol con ese nombre.')

        else:
            try:
                with transaction.atomic():
                    rol = Rol.objects.create(nombrerol=nombre)

                    permisos_validos = Permiso.objects.filter(
                        pk__in=permisos_seleccionados
                    )

                    RolPermiso.objects.bulk_create([
                        RolPermiso(
                            idrol=rol,
                            idpermiso=permiso
                        )
                        for permiso in permisos_validos
                    ])

                messages.success(
                    request,
                    'Rol creado y permisos asignados correctamente.'
                )
                return redirect('rol_lista')

            except Exception:
                logger.exception(
                    'Error al crear el rol "%s" y asignar sus permisos.',
                    nombre
                )
                messages.error(
                    request,
                    'No se pudo crear el rol. Revisa el error en la consola del servidor.'
                )

    return render(
        request,
        'panel/administracion/roles/rolCrear.html',
        {
            'titulo': 'Nuevo rol',
            'rol': None,
            'permisos': permisos,
            'permisos_seleccionados': permisos_seleccionados,
        }
    )


def rol_editar(request, idrol):
    rol = get_object_or_404(Rol, pk=idrol)

    # Obtener todos los permisos disponibles
    permisos = Permiso.objects.all().order_by('modulo', 'accion')

    # Obtener los permisos que ya tiene asignados el rol
    permisos_seleccionados = list(
        RolPermiso.objects.filter(
            idrol=rol
        ).values_list('idpermiso_id', flat=True)
    )

    if request.method == 'POST':
        nombre = request.POST.get('nombrerol', '').strip()

        # Obtener las casillas seleccionadas
        ids_permisos = []

        for valor in request.POST.getlist('permisos'):
            try:
                ids_permisos.append(int(valor))
            except (ValueError, TypeError):
                continue

        # Validar el nombre
        if not nombre:
            messages.error(request, 'El nombre del rol es obligatorio.')

        elif len(nombre) > 50:
            messages.error(
                request,
                'El nombre del rol no puede superar los 50 caracteres.'
            )

        elif Rol.objects.filter(
            nombrerol__iexact=nombre
        ).exclude(pk=rol.pk).exists():
            messages.error(request, 'Ya existe un rol con ese nombre.')

        else:
            try:
                with transaction.atomic():
                    # Actualizar el nombre
                    rol.nombrerol = nombre
                    rol.save(update_fields=['nombrerol'])

                    # Guardar los permisos seleccionados
                    permisos_validos = Permiso.objects.filter(
                        pk__in=ids_permisos
                    )

                    # Quitar las asignaciones anteriores
                    RolPermiso.objects.filter(idrol=rol).delete()

                    # Crear las nuevas asignaciones
                    RolPermiso.objects.bulk_create([
                        RolPermiso(
                            idrol=rol,
                            idpermiso=permiso
                        )
                        for permiso in permisos_validos
                    ])

                messages.success(
                    request,
                    'Rol y permisos actualizados correctamente.'
                )
                return redirect('rol_lista')

            except Exception:
                logger.exception(
                    'Error al editar el rol "%s".',
                    nombre
                )
                messages.error(
                    request,
                    'No se pudo actualizar el rol y sus permisos. '
                    'Revisa la consola del servidor.'
                )

    return render(
        request,
        'panel/administracion/roles/rolEditar.html',
        {
            'titulo': 'Editar rol',
            'rol': rol,
            'permisos': permisos,
            'permisos_seleccionados': permisos_seleccionados,
        }
    )

@require_POST
def rol_eliminar(request, idrol):
    rol = get_object_or_404(Rol, pk=idrol)
    nombre = rol.nombrerol or f"Rol {rol.pk}"

    try:
        rol.delete()
        messages.success(
            request,
            f'El rol "{nombre}" se eliminó correctamente.'
        )
    except ProtectedError:
        messages.error(
            request,
            f'No se puede eliminar el rol "{nombre}" porque está '
            'asignado a empleados u otros registros relacionados.'
        )
    except Exception:
        messages.error(
            request,
            f'No se pudo eliminar el rol "{nombre}".'
        )

    return redirect('rol_lista')

def rol_permisos(request, idrol):
    rol = get_object_or_404(Rol, pk=idrol)

    permisos = Permiso.objects.all().order_by('modulo', 'accion')

    permisos_asignados = RolPermiso.objects.filter(
        idrol=rol
    ).values_list('idpermiso_id', flat=True)

    return render(
        request,
        'panel/administracion/roles/rolPermisos.html',
        {
            'rol': rol,
            'permisos': permisos,
            'permisos_asignados': list(permisos_asignados),
        },
    )


 
# ============================================================
# 10. PANEL - REPORTES
# ============================================================

# ------------------------------------------------------------
# 10.1 Constantes de diseño
# ------------------------------------------------------------

BI_OSCURO = "17181D"
BI_ROJO = "D93632"
BI_GRIS = "6B7280"
BI_PALETA = [
    "D93632", "3B82F6", "10B981", "F59E0B",
    "8B5CF6", "F97316", "06B6D4", "EC4899",
]

BI_FORMATO_XLSX = {
    "moneda": '"$" #,##0.00',
    "entero": "#,##0",
    "pct": '0.0"%"',
    "pct_signo": '+0.0"%";-0.0"%"',
    "fecha": "DD/MM/YYYY",
}
BI_NUMERICOS = {"moneda", "entero", "pct", "pct_signo"}

BI_DIAS_SEMANA = [
    "Lunes", "Martes", "Miércoles", "Jueves",
    "Viernes", "Sábado", "Domingo",
]


# ------------------------------------------------------------
# 10.2 Utilidades de fechas y formato
# ------------------------------------------------------------

def bi_fechas(request):
    """Lee y valida el rango de fechas enviado por el usuario."""
    hoy = timezone.localdate()
    inicio_default = hoy - timedelta(days=29)

    inicio = parse_date(request.GET.get("fecha_inicio", "")) or inicio_default
    fin = parse_date(request.GET.get("fecha_fin", "")) or hoy

    if inicio > fin:
        inicio, fin = fin, inicio

    return inicio, fin


def bi_contexto(inicio, fin):
    return {
        "fecha_inicio": inicio,
        "fecha_fin": fin,
        "periodo": f"{inicio:%d/%m/%Y} al {fin:%d/%m/%Y}",
    }


def bi_periodo_anterior(inicio, fin):
    """Período de igual duración inmediatamente anterior al actual."""
    duracion = (fin - inicio).days + 1
    fin_anterior = inicio - timedelta(days=1)
    inicio_anterior = fin_anterior - timedelta(days=duracion - 1)
    return inicio_anterior, fin_anterior, duracion


def bi_variacion(actual, anterior):
    """Variación porcentual; None si no hay base de comparación."""
    if anterior and anterior > 0:
        return (actual - anterior) / anterior * Decimal("100")
    return None


def bi_cero(valor):
    return valor if valor is not None else Decimal("0.00")


def bi_pct(parte, total):
    return float(parte / total * 100) if total else 0.0


def bi_numero(valor, decimales=0):
    """Formato argentino: 1.234,56"""
    texto = f"{valor:,.{decimales}f}"
    return texto.replace(",", "X").replace(".", ",").replace("X", ".")


def bi_moneda(valor):
    return f"$ {bi_numero(bi_cero(valor), 2)}"


def bi_formatear(valor, formato=None):
    if valor is None or valor == "":
        return "—"
    if formato == "moneda":
        return bi_moneda(valor)
    if formato == "entero":
        return bi_numero(valor, 0)
    if formato == "pct":
        return f"{bi_numero(valor, 1)}%"
    if formato == "pct_signo":
        return f"{'+' if valor > 0 else ''}{bi_numero(valor, 1)}%"
    if formato == "fecha":
        return valor.strftime("%d/%m/%Y")
    return str(valor)


def bi_valor_excel(valor):
    if valor is None:
        return "—"
    if isinstance(valor, Decimal):
        return float(valor)
    return valor


# ------------------------------------------------------------
# 10.3 Consultas base
# ------------------------------------------------------------

def bi_pedidos_validos():
    """Pedidos no cancelados (base de las ventas)."""
    return (
        Pedido.objects
        .exclude(idestadopedido__descripcion="Cancelado")
        .filter(total__isnull=False)
    )


def bi_pedidos_periodo(inicio, fin):
    return bi_pedidos_validos().filter(
        fecha_creacion__date__gte=inicio,
        fecha_creacion__date__lte=fin,
    )


def bi_pedidos_pagados_validos():
    """Pedidos cobrados y no cancelados (base del dinero recaudado)."""
    return (
        Pedido.objects
        .filter(pagado=True, total__isnull=False)
        .exclude(idestadopedido__descripcion="Cancelado")
    )


def bi_pedidos_pagados_periodo(inicio, fin):
    return bi_pedidos_pagados_validos().filter(
        fecha_creacion__date__gte=inicio,
        fecha_creacion__date__lte=fin,
    )


# ------------------------------------------------------------
# 10.4 Cálculos de cada reporte
# ------------------------------------------------------------

def bi_resumen(inicio, fin):
    pedidos = bi_pedidos_periodo(inicio, fin)

    indicadores = pedidos.aggregate(
        total_vendido=Sum("total"),
        cantidad_pedidos=Count("idpedido"),
        ticket_promedio=Avg("total"),
    )
    indicadores["total_vendido"] = bi_cero(indicadores["total_vendido"])
    indicadores["cantidad_pedidos"] = indicadores["cantidad_pedidos"] or 0
    indicadores["ticket_promedio"] = bi_cero(indicadores["ticket_promedio"])

    ventas_diarias = list(
        pedidos.annotate(dia=TruncDate("fecha_creacion"))
        .values("dia")
        .annotate(total=Sum("total"), pedidos=Count("idpedido"))
        .order_by("dia")
    )

    maximo = max(
        (bi_cero(venta["total"]) for venta in ventas_diarias),
        default=Decimal("0.00"),
    )

    for venta in ventas_diarias:
        venta["total"] = bi_cero(venta["total"])
        venta["pedidos"] = venta["pedidos"] or 0
        venta["ticket_promedio"] = (
            venta["total"] / venta["pedidos"]
            if venta["pedidos"]
            else Decimal("0.00")
        )
        venta["porcentaje"] = (
            float(venta["total"] / maximo * 100) if maximo > 0 else 0
        )

    return indicadores, ventas_diarias


def bi_metricas_adicionales(inicio, fin, indicadores, ventas_diarias):
    inicio_anterior, fin_anterior, duracion = bi_periodo_anterior(inicio, fin)

    total_anterior = bi_cero(
        bi_pedidos_periodo(inicio_anterior, fin_anterior)
        .aggregate(total=Sum("total"))["total"]
    )
    total_actual = indicadores["total_vendido"]

    pedidos_por_estado = list(
        Pedido.objects.filter(
            fecha_creacion__date__gte=inicio,
            fecha_creacion__date__lte=fin,
        )
        .values("idestadopedido__descripcion")
        .annotate(cantidad=Count("idpedido"))
        .order_by("idestadopedido__descripcion")
    )

    return {
        "total_anterior": total_anterior,
        "inicio_anterior": inicio_anterior,
        "fin_anterior": fin_anterior,
        "variacion_porcentual": bi_variacion(total_actual, total_anterior),
        "promedio_diario": (
            total_actual / Decimal(duracion) if duracion > 0 else Decimal("0.00")
        ),
        "mejor_dia": max(
            ventas_diarias, key=lambda venta: venta["total"], default=None
        ),
        "pedidos_por_estado": pedidos_por_estado,
    }


def bi_top_productos(inicio, fin):
    importe_linea = ExpressionWrapper(
        F("cantidad") * F("preciounitario"),
        output_field=DecimalField(max_digits=14, decimal_places=2),
    )

    return list(
        DetallePedido.objects.filter(
            idpedido__in=bi_pedidos_periodo(inicio, fin),
            idproducto__isnull=False,
        )
        .values("idproducto_id", "idproducto__nombre")
        .annotate(unidades=Sum("cantidad"), ingresos=Sum(importe_linea))
        .order_by("-unidades", "-ingresos", "idproducto__nombre")[:10]
    )


def bi_totales_productos(productos):
    unidades = sum((p["unidades"] or 0) for p in productos)
    ingresos = sum(
        (bi_cero(p["ingresos"]) for p in productos), Decimal("0.00")
    )
    return unidades, ingresos


def bi_datos_dinero(inicio, fin):
    pedidos = bi_pedidos_pagados_periodo(inicio, fin)

    indicadores = pedidos.aggregate(
        importe_total=Sum("total"),
        cantidad_pedidos=Count("idpedido"),
        ticket_promedio=Avg("total"),
    )
    importe_total = bi_cero(indicadores["importe_total"])
    cantidad_pedidos = indicadores["cantidad_pedidos"] or 0
    ticket_promedio = bi_cero(indicadores["ticket_promedio"])

    inicio_anterior, fin_anterior, duracion = bi_periodo_anterior(inicio, fin)
    promedio_diario = (
        importe_total / Decimal(duracion) if duracion > 0 else Decimal("0.00")
    )

    # --- Serie diaria (incluye días sin cobros) ---
    consulta_diaria = list(
        pedidos.annotate(fecha=TruncDate("fecha_creacion"))
        .values("fecha")
        .annotate(importe=Sum("total"), cantidad=Count("idpedido"))
        .order_by("fecha")
    )
    por_fecha = {
        fila["fecha"]: fila for fila in consulta_diaria if fila["fecha"]
    }
    maximo = max(
        (bi_cero(fila["importe"]) for fila in consulta_diaria),
        default=Decimal("0.00"),
    )

    dias = []
    fecha = inicio
    while fecha <= fin:
        fila = por_fecha.get(fecha, {})
        importe = bi_cero(fila.get("importe"))
        dias.append({
            "fecha": fecha,
            "importe": importe,
            "cantidad": fila.get("cantidad") or 0,
            "porcentaje": float(importe / maximo * 100) if maximo > 0 else 0,
        })
        fecha += timedelta(days=1)

    dias_con_datos = [dia for dia in dias if dia["cantidad"] > 0]
    mejor_dia = max(dias_con_datos, key=lambda d: d["importe"], default=None)
    peor_dia = min(dias_con_datos, key=lambda d: d["importe"], default=None)

    # --- Por día de la semana ---
    agrupados = {
        indice: {"dia": nombre, "importe": Decimal("0.00"), "cantidad": 0}
        for indice, nombre in enumerate(BI_DIAS_SEMANA)
    }
    for dia in dias:
        indice = dia["fecha"].weekday()
        agrupados[indice]["importe"] += dia["importe"]
        agrupados[indice]["cantidad"] += dia["cantidad"]

    maximo_semana = max(
        (fila["importe"] for fila in agrupados.values()),
        default=Decimal("0.00"),
    )
    dias_semana = []
    for fila in agrupados.values():
        fila["porcentaje"] = (
            float(fila["importe"] / maximo_semana * 100)
            if maximo_semana > 0
            else 0
        )
        dias_semana.append(fila)

    mejor_dia_semana = max(
        dias_semana, key=lambda d: d["importe"], default=None
    )

    # --- Por medio de pago ---
    consulta_medios = list(
        pedidos.values("idmediopago__nombremetodo")
        .annotate(cantidad=Count("idpedido"), importe=Sum("total"))
        .order_by("-importe")
    )
    medios_pago = []
    for medio in consulta_medios:
        importe = bi_cero(medio["importe"])
        medios_pago.append({
            "nombre": medio["idmediopago__nombremetodo"] or "Sin medio de pago",
            "cantidad": medio["cantidad"],
            "importe": importe,
            "porcentaje": bi_pct(importe, importe_total),
        })

    # --- Comparación con el período anterior ---
    importe_anterior = bi_cero(
        bi_pedidos_pagados_periodo(inicio_anterior, fin_anterior)
        .aggregate(total=Sum("total"))["total"]
    )

    return {
        "importe_total": importe_total,
        "cantidad_pedidos": cantidad_pedidos,
        "ticket_promedio": ticket_promedio,
        "promedio_diario": promedio_diario,
        "dias": dias,
        "mejor_dia": mejor_dia,
        "peor_dia": peor_dia,
        "dias_semana": dias_semana,
        "mejor_dia_semana": mejor_dia_semana,
        "medios_pago": medios_pago,
        "variacion_porcentual": bi_variacion(importe_total, importe_anterior),
        "importe_anterior": importe_anterior,
        "inicio_anterior": inicio_anterior,
        "fin_anterior": fin_anterior,
    }


# ------------------------------------------------------------
# 10.5 Definición de los reportes exportables
#      (misma estructura para Excel y PDF)
# ------------------------------------------------------------

def bi_reporte_ventas(inicio, fin):
    indicadores, ventas = bi_resumen(inicio, fin)
    extra = bi_metricas_adicionales(inicio, fin, indicadores, ventas)
    total = indicadores["total_vendido"]
    mejor = extra["mejor_dia"]
    anterior = f"{extra['inicio_anterior']:%d/%m} al {extra['fin_anterior']:%d/%m}"

    estados = [
        (fila["idestadopedido__descripcion"] or "Sin estado", fila["cantidad"])
        for fila in extra["pedidos_por_estado"]
    ]
    total_estados = sum(cantidad for _, cantidad in estados)

    return {
        "titulo": "Resumen de ventas",
        "periodo": f"{inicio:%d/%m/%Y} al {fin:%d/%m/%Y}",
        "kpis": [
            {"titulo": "Total vendido", "valor": total, "formato": "moneda",
             "detalle": "Pedidos no cancelados"},
            {"titulo": "Pedidos", "valor": indicadores["cantidad_pedidos"],
             "formato": "entero", "detalle": "Operaciones del período"},
            {"titulo": "Ticket promedio", "valor": indicadores["ticket_promedio"],
             "formato": "moneda", "detalle": "Importe por pedido"},
            {"titulo": "Promedio diario", "valor": extra["promedio_diario"],
             "formato": "moneda", "detalle": "Ventas por día"},
            {"titulo": "Variación", "valor": extra["variacion_porcentual"],
             "formato": "pct_signo", "detalle": f"vs. {anterior}"},
            {"titulo": "Mejor día", "valor": mejor["dia"] if mejor else None,
             "formato": "fecha",
             "detalle": bi_moneda(mejor["total"]) if mejor else "Sin ventas"},
        ],
        "graficos": [
            {"tipo": "linea", "titulo": "Evolución de ventas por día",
             "etiquetas": [f"{v['dia']:%d/%m}" for v in ventas],
             "valores": [v["total"] for v in ventas]},
            {"tipo": "barras", "titulo": "Pedidos por día",
             "etiquetas": [f"{v['dia']:%d/%m}" for v in ventas],
             "valores": [v["pedidos"] for v in ventas]},
            {"tipo": "torta", "titulo": "Pedidos por estado (incluye cancelados)",
             "etiquetas": [e for e, _ in estados],
             "valores": [c for _, c in estados]},
        ],
        "tablas": [
            {
                "titulo": "Detalle diario de ventas",
                "encabezados": ["Fecha", "Pedidos", "Ticket promedio",
                                "Total vendido", "% del período"],
                "formatos": ["fecha", "entero", "moneda", "moneda", "pct"],
                "filas": [
                    [v["dia"], v["pedidos"], v["ticket_promedio"], v["total"],
                     bi_pct(v["total"], total)]
                    for v in ventas
                ],
            },
            {
                "titulo": "Pedidos por estado",
                "encabezados": ["Estado", "Pedidos", "% del total"],
                "formatos": [None, "entero", "pct"],
                "filas": [[e, c, bi_pct(c, total_estados)] for e, c in estados],
            },
        ],
    }


def bi_reporte_productos(inicio, fin):
    productos = bi_top_productos(inicio, fin)
    unidades_total, ingresos_total = bi_totales_productos(productos)

    lider = productos[0] if productos else None
    mayor_ingreso = max(
        productos, key=lambda p: bi_cero(p["ingresos"]), default=None
    )
    precio_promedio = (
        ingresos_total / unidades_total if unidades_total else Decimal("0.00")
    )
    nombres = [p["idproducto__nombre"] or "Sin nombre" for p in productos]

    return {
        "titulo": "Productos más vendidos",
        "periodo": f"{inicio:%d/%m/%Y} al {fin:%d/%m/%Y}",
        "kpis": [
            {"titulo": "Unidades (Top 10)", "valor": unidades_total,
             "formato": "entero", "detalle": "Suma del ranking"},
            {"titulo": "Ingresos (Top 10)", "valor": ingresos_total,
             "formato": "moneda", "detalle": "Suma del ranking"},
            {"titulo": "Precio promedio", "valor": precio_promedio,
             "formato": "moneda", "detalle": "Ingreso por unidad"},
            {"titulo": "Más vendido",
             "valor": lider["idproducto__nombre"] if lider else None,
             "formato": None,
             "detalle": f"{lider['unidades']} unidades" if lider else "Sin datos"},
            {"titulo": "Peso del líder",
             "valor": bi_pct(lider["unidades"], unidades_total) if lider else None,
             "formato": "pct", "detalle": "% de unidades del Top 10"},
            {"titulo": "Mayor ingreso",
             "valor": mayor_ingreso["idproducto__nombre"] if mayor_ingreso else None,
             "formato": None,
             "detalle": bi_moneda(mayor_ingreso["ingresos"]) if mayor_ingreso else "Sin datos"},
        ],
        "graficos": [
            {"tipo": "barras_h", "titulo": "Unidades vendidas por producto",
             "etiquetas": nombres,
             "valores": [p["unidades"] or 0 for p in productos]},
            {"tipo": "barras_h", "titulo": "Ingresos por producto",
             "etiquetas": nombres,
             "valores": [bi_cero(p["ingresos"]) for p in productos]},
        ],
        "tablas": [
            {
                "titulo": "Ranking detallado (participación sobre el Top 10)",
                "encabezados": ["#", "Producto", "Unidades", "% unidades",
                                "Ingresos", "% ingresos", "Precio promedio"],
                "formatos": ["entero", None, "entero", "pct",
                             "moneda", "pct", "moneda"],
                "filas": [
                    [
                        posicion,
                        nombre,
                        p["unidades"] or 0,
                        bi_pct(p["unidades"] or 0, unidades_total),
                        bi_cero(p["ingresos"]),
                        bi_pct(bi_cero(p["ingresos"]), ingresos_total),
                        (
                            bi_cero(p["ingresos"]) / p["unidades"]
                            if p["unidades"]
                            else Decimal("0.00")
                        ),
                    ]
                    for posicion, (nombre, p) in enumerate(
                        zip(nombres, productos), start=1
                    )
                ],
            }
        ],
    }


def bi_reporte_dinero(inicio, fin):
    d = bi_datos_dinero(inicio, fin)
    total = d["importe_total"]
    mejor = d["mejor_dia"]
    anterior = f"{d['inicio_anterior']:%d/%m} al {d['fin_anterior']:%d/%m}"

    return {
        "titulo": "Dinero recaudado",
        "periodo": f"{inicio:%d/%m/%Y} al {fin:%d/%m/%Y}",
        "kpis": [
            {"titulo": "Recaudación total", "valor": total,
             "formato": "moneda", "detalle": "Pedidos pagados"},
            {"titulo": "Pedidos cobrados", "valor": d["cantidad_pedidos"],
             "formato": "entero", "detalle": "Operaciones pagadas"},
            {"titulo": "Ticket promedio", "valor": d["ticket_promedio"],
             "formato": "moneda", "detalle": "Importe por pedido"},
            {"titulo": "Promedio diario", "valor": d["promedio_diario"],
             "formato": "moneda", "detalle": "Cobrado por día"},
            {"titulo": "Variación", "valor": d["variacion_porcentual"],
             "formato": "pct_signo", "detalle": f"vs. {anterior}"},
            {"titulo": "Mejor día", "valor": mejor["fecha"] if mejor else None,
             "formato": "fecha",
             "detalle": bi_moneda(mejor["importe"]) if mejor else "Sin cobros"},
        ],
        "graficos": [
            {"tipo": "linea", "titulo": "Recaudación por día",
             "etiquetas": [f"{x['fecha']:%d/%m}" for x in d["dias"]],
             "valores": [x["importe"] for x in d["dias"]]},
            {"tipo": "torta", "titulo": "Recaudación por medio de pago",
             "etiquetas": [m["nombre"] for m in d["medios_pago"]],
             "valores": [m["importe"] for m in d["medios_pago"]]},
            {"tipo": "barras", "titulo": "Recaudación por día de la semana",
             "etiquetas": [s["dia"] for s in d["dias_semana"]],
             "valores": [s["importe"] for s in d["dias_semana"]]},
        ],
        "tablas": [
            {
                "titulo": "Detalle por medio de pago",
                "encabezados": ["Medio de pago", "Pedidos", "Importe", "% del total"],
                "formatos": [None, "entero", "moneda", "pct"],
                "filas": [
                    [m["nombre"], m["cantidad"], m["importe"], m["porcentaje"]]
                    for m in d["medios_pago"]
                ],
            },
            {
                "titulo": "Rendimiento por día de la semana",
                "encabezados": ["Día", "Pedidos", "Importe", "% del total"],
                "formatos": [None, "entero", "moneda", "pct"],
                "filas": [
                    [s["dia"], s["cantidad"], s["importe"],
                     bi_pct(s["importe"], total)]
                    for s in d["dias_semana"]
                ],
            },
            {
                "titulo": "Detalle diario",
                "encabezados": ["Fecha", "Pedidos cobrados", "Recaudación"],
                "formatos": ["fecha", "entero", "moneda"],
                "filas": [
                    [x["fecha"], x["cantidad"], x["importe"]]
                    for x in d["dias"]
                    if x["cantidad"]
                ],
            },
        ],
    }


# ------------------------------------------------------------
# 10.6 Exportación a Excel
# ------------------------------------------------------------

def bi_celda(hoja, rango, valor, fuente, relleno=None,
             alineacion=None, formato=None):
    hoja.merge_cells(rango)
    celda = hoja[rango.split(":")[0]]
    celda.value = valor
    celda.font = fuente
    if relleno:
        celda.fill = PatternFill("solid", fgColor=relleno)
    celda.alignment = alineacion or Alignment(
        horizontal="left", vertical="center", indent=1
    )
    if formato:
        celda.number_format = formato
    return celda


def bi_grafico_excel(grafico, datos, indice):
    etiquetas = grafico["etiquetas"]
    valores = grafico["valores"]
    cantidad = len(etiquetas)
    if cantidad == 0 or not any(valores):
        return None

    columna = indice * 3 + 1
    datos.cell(row=1, column=columna, value="Etiqueta")
    datos.cell(row=1, column=columna + 1, value=grafico["titulo"])
    for numero, (etiqueta, valor) in enumerate(zip(etiquetas, valores), start=2):
        datos.cell(row=numero, column=columna, value=str(etiqueta))
        datos.cell(row=numero, column=columna + 1, value=float(valor))

    categorias = Reference(datos, min_col=columna, min_row=2, max_row=cantidad + 1)
    serie = Reference(datos, min_col=columna + 1, min_row=1, max_row=cantidad + 1)
    tipo = grafico["tipo"]

    if tipo == "torta":
        objeto = PieChart()
        objeto.add_data(serie, titles_from_data=True)
        objeto.set_categories(categorias)
        objeto.dataLabels = DataLabelList()
        objeto.dataLabels.showPercent = True
        objeto.dataLabels.showVal = False
        objeto.dataLabels.showCatName = False
        objeto.dataLabels.showSerName = False
        objeto.dataLabels.showLeaderLines = False
        for posicion in range(cantidad):
            objeto.series[0].dPt.append(
                DataPoint(
                    idx=posicion,
                    spPr=GraphicalProperties(
                        solidFill=BI_PALETA[posicion % len(BI_PALETA)]
                    ),
                )
            )
    else:
        objeto = LineChart() if tipo == "linea" else BarChart()
        if tipo != "linea":
            objeto.type = "bar" if tipo == "barras_h" else "col"
            objeto.gapWidth = 60
        objeto.add_data(serie, titles_from_data=True)
        objeto.set_categories(categorias)
        objeto.legend = None
        objeto.x_axis.delete = False
        objeto.y_axis.delete = False
        objeto.y_axis.number_format = "#,##0"

        s = objeto.series[0]
        if tipo == "linea":
            s.smooth = True
            s.graphicalProperties.line.solidFill = BI_ROJO
            s.graphicalProperties.line.width = 28575
            s.marker = Marker(
                symbol="circle", size=6,
                spPr=GraphicalProperties(solidFill=BI_ROJO),
            )
        else:
            s.graphicalProperties.solidFill = BI_ROJO

        if tipo == "barras_h":
            objeto.x_axis.scaling.orientation = "maxMin"  # el #1 arriba
            objeto.y_axis.crosses = "max"                 # eje de valores abajo

    objeto.title = grafico["titulo"]
    objeto.width = 16
    objeto.height = 7.5
    return objeto


def bi_exportar_excel(reporte, nombre_archivo):
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Reporte"
    hoja.sheet_view.showGridLines = False
    hoja.sheet_properties.tabColor = BI_ROJO
    datos = libro.create_sheet("Datos")

    COLUMNAS = 12
    ultima = get_column_letter(COLUMNAS)
    for numero in range(1, COLUMNAS + 1):
        hoja.column_dimensions[get_column_letter(numero)].width = 14

    # ---------- Encabezado ----------
    for fila in (1, 2, 3):
        for numero in range(1, COLUMNAS + 1):
            hoja.cell(row=fila, column=numero).fill = PatternFill(
                "solid", fgColor=BI_OSCURO
            )
    bi_celda(hoja, f"A1:{ultima}1", "ZONA BURGER",
             Font(bold=True, size=20, color="FFFFFF"))
    bi_celda(hoja, f"A2:{ultima}2", reporte["titulo"],
             Font(bold=True, size=13, color="FF6B66"))
    bi_celda(
        hoja, f"A3:{ultima}3",
        f"Período: {reporte['periodo']}   |   "
        f"Generado el {timezone.localdate():%d/%m/%Y}",
        Font(size=10, color="9CA3AF"),
    )
    hoja.row_dimensions[1].height = 34
    hoja.row_dimensions[2].height = 22
    hoja.row_dimensions[3].height = 20

    # ---------- Tarjetas de KPI ----------
    for indice, kpi in enumerate(reporte["kpis"][:6]):
        c1 = get_column_letter(indice * 2 + 1)
        c2 = get_column_letter(indice * 2 + 2)
        formato = kpi.get("formato")
        valor = bi_valor_excel(kpi["valor"])

        titulo_kpi = bi_celda(
            hoja, f"{c1}5:{c2}5", kpi["titulo"].upper(),
            Font(bold=True, size=8, color="9CA3AF"), BI_OSCURO,
        )
        numero_kpi = bi_celda(
            hoja, f"{c1}6:{c2}6", valor,
            Font(bold=True, size=18, color="FFFFFF"), BI_OSCURO,
            formato=BI_FORMATO_XLSX.get(formato),
        )
        detalle_kpi = bi_celda(
            hoja, f"{c1}7:{c2}7", kpi.get("detalle", ""),
            Font(size=8, color="9CA3AF"), BI_OSCURO,
        )
        titulo_kpi.border = Border(
            left=Side(style="thick", color="FFFFFF"),
            top=Side(style="thick", color=BI_ROJO),
        )
        numero_kpi.border = Border(left=Side(style="thick", color="FFFFFF"))
        detalle_kpi.border = Border(left=Side(style="thick", color="FFFFFF"))
        hoja[f"{c2}5"].border = Border(top=Side(style="thick", color=BI_ROJO))
    hoja.row_dimensions[5].height = 20
    hoja.row_dimensions[6].height = 36
    hoja.row_dimensions[7].height = 18

    # ---------- Gráficos (2 por fila) ----------
    fila = 9
    graficos = reporte.get("graficos", [])
    for indice, grafico in enumerate(graficos):
        objeto = bi_grafico_excel(grafico, datos, indice)
        if objeto is None:
            continue
        columna = "A" if indice % 2 == 0 else "G"
        hoja.add_chart(objeto, f"{columna}{fila + (indice // 2) * 17}")
    fila += ((len(graficos) + 1) // 2) * 17 + 1

    # ---------- Tablas ----------
    for tabla in reporte.get("tablas", []):
        encabezados = tabla["encabezados"]
        formatos = tabla.get("formatos") or [None] * len(encabezados)
        ancho = max(len(encabezados), 2)

        bi_celda(
            hoja, f"A{fila}:{get_column_letter(ancho)}{fila}",
            tabla["titulo"].upper(),
            Font(bold=True, size=11, color="FFFFFF"), BI_OSCURO,
        )
        hoja.row_dimensions[fila].height = 22
        fila += 1

        for numero, encabezado in enumerate(encabezados, start=1):
            celda = hoja.cell(row=fila, column=numero, value=encabezado)
            celda.font = Font(bold=True, color="FFFFFF")
            celda.fill = PatternFill("solid", fgColor=BI_ROJO)
            celda.alignment = Alignment(
                horizontal="right" if formatos[numero - 1] in BI_NUMERICOS else "left",
                vertical="center",
                wrap_text=True,
            )
        hoja.row_dimensions[fila].height = 22
        fila += 1

        borde = Border(bottom=Side(style="thin", color="E5E7EB"))
        for posicion, datos_fila in enumerate(tabla["filas"]):
            for numero, valor in enumerate(datos_fila, start=1):
                formato = formatos[numero - 1]
                celda = hoja.cell(row=fila, column=numero, value=bi_valor_excel(valor))
                if formato in BI_FORMATO_XLSX:
                    celda.number_format = BI_FORMATO_XLSX[formato]
                celda.alignment = Alignment(
                    horizontal="right" if formato in BI_NUMERICOS else "left",
                    vertical="center",
                )
                celda.border = borde
                if posicion % 2:
                    celda.fill = PatternFill("solid", fgColor="F3F4F6")
            fila += 1
        fila += 2

    # ---------- Impresión ----------
    datos.sheet_state = "hidden"
    hoja.page_setup.orientation = "landscape"
    hoja.page_setup.fitToWidth = 1
    hoja.page_setup.fitToHeight = 0
    hoja.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)

    buffer = BytesIO()
    libro.save(buffer)
    respuesta = HttpResponse(
        buffer.getvalue(),
        content_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
    )
    respuesta["Content-Disposition"] = f'attachment; filename="{nombre_archivo}"'
    return respuesta


# ------------------------------------------------------------
# 10.7 Exportación a PDF
# ------------------------------------------------------------

def bi_dibujo_grafico(grafico, ancho, alto):
    dibujo = Drawing(ancho, alto)
    rojo = colors.HexColor("#" + BI_ROJO)
    etiquetas = [str(e) for e in grafico["etiquetas"]]
    valores = [float(v) for v in grafico["valores"]]

    if not valores or not any(valores):
        dibujo.add(String(
            ancho / 2, alto / 2, "Sin datos para el período",
            textAnchor="middle", fontSize=9,
            fillColor=colors.HexColor("#" + BI_GRIS),
        ))
        return dibujo

    tipo = grafico["tipo"]

    if tipo == "torta":
        torta = Pie()
        torta.width = torta.height = alto - 50
        torta.x = ancho * 0.30
        torta.y = 20
        torta.data = valores
        torta.labels = [e[:22] for e in etiquetas]
        torta.sideLabels = True
        torta.slices.strokeColor = colors.white
        torta.slices.strokeWidth = 1
        torta.slices.fontSize = 7
        for posicion in range(len(valores)):
            torta.slices[posicion].fillColor = colors.HexColor(
                "#" + BI_PALETA[posicion % len(BI_PALETA)]
            )
        dibujo.add(torta)
        return dibujo

    if tipo == "barras_h":
        graf = HorizontalBarChart()
        graf.x, graf.y = 115, 10
        graf.width, graf.height = ancho - 135, alto - 20
        graf.data = [list(reversed(valores))]
        graf.categoryAxis.categoryNames = [e[:20] for e in reversed(etiquetas)]
        graf.categoryAxis.labels.fontSize = 7
        graf.valueAxis.labels.fontSize = 7
        graf.valueAxis.valueMin = 0
        graf.bars[0].fillColor = rojo
        graf.bars[0].strokeColor = None
        dibujo.add(graf)
        return dibujo

    graf = HorizontalLineChart() if tipo == "linea" else VerticalBarChart()
    graf.x, graf.y = 50, 35
    graf.width, graf.height = ancho - 70, alto - 50
    graf.data = [valores]
    paso = max(1, len(etiquetas) // 10)
    graf.categoryAxis.categoryNames = [
        e if i % paso == 0 else "" for i, e in enumerate(etiquetas)
    ]
    graf.categoryAxis.labels.fontSize = 7
    graf.categoryAxis.labels.angle = 30
    graf.categoryAxis.labels.boxAnchor = "ne"
    graf.valueAxis.labels.fontSize = 7
    graf.valueAxis.valueMin = 0
    graf.valueAxis.visibleGrid = True
    graf.valueAxis.gridStrokeColor = colors.HexColor("#E5E7EB")
    if tipo == "linea":
        graf.lines[0].strokeColor = rojo
        graf.lines[0].strokeWidth = 2
        graf.lines[0].symbol = makeMarker("Circle")
    else:
        graf.bars[0].fillColor = rojo
        graf.bars[0].strokeColor = None
    dibujo.add(graf)
    return dibujo


def bi_marco_pdf(titulo, periodo):
    """Cabecera y pie que se repiten en todas las páginas."""
    def dibujar(canvas, documento):
        ancho, alto = landscape(A4)
        canvas.saveState()
        canvas.setFillColor(colors.HexColor("#" + BI_OSCURO))
        canvas.rect(0, alto - 22 * mm, ancho, 22 * mm, stroke=0, fill=1)
        canvas.setFillColor(colors.HexColor("#" + BI_ROJO))
        canvas.rect(0, alto - 22 * mm, 6 * mm, 22 * mm, stroke=0, fill=1)
        canvas.setFillColor(colors.white)
        canvas.setFont("Helvetica-Bold", 16)
        canvas.drawString(15 * mm, alto - 10 * mm, "ZONA BURGER")
        canvas.setFont("Helvetica", 10)
        canvas.drawString(15 * mm, alto - 16.5 * mm, titulo)
        canvas.drawRightString(ancho - 15 * mm, alto - 10 * mm, f"Período: {periodo}")
        canvas.setFont("Helvetica", 8)
        canvas.drawRightString(
            ancho - 15 * mm, alto - 16.5 * mm,
            f"Generado el {timezone.localdate():%d/%m/%Y}",
        )
        canvas.setFillColor(colors.HexColor("#" + BI_GRIS))
        canvas.drawString(15 * mm, 8 * mm, "Zona Burger · Reporte de gestión")
        canvas.drawRightString(ancho - 15 * mm, 8 * mm, f"Página {documento.page}")
        canvas.restoreState()
    return dibujar


def bi_exportar_pdf(reporte, nombre_archivo):
    respuesta = HttpResponse(content_type="application/pdf")
    respuesta["Content-Disposition"] = f'attachment; filename="{nombre_archivo}"'

    pagina = landscape(A4)
    ancho_total = pagina[0] - 30 * mm
    documento = SimpleDocTemplate(
        respuesta, pagesize=pagina,
        leftMargin=15 * mm, rightMargin=15 * mm,
        topMargin=30 * mm, bottomMargin=16 * mm,
    )

    oscuro = colors.HexColor("#" + BI_OSCURO)
    rojo = colors.HexColor("#" + BI_ROJO)
    borde = colors.HexColor("#E5E7EB")
    gris = colors.HexColor("#9CA3AF")
    base = getSampleStyleSheet()["Normal"]

    est_kpi_titulo = ParagraphStyle("kpiT", parent=base, fontSize=7,
                                    leading=9, textColor=gris)
    est_kpi_valor = ParagraphStyle("kpiV", parent=base, fontName="Helvetica-Bold",
                                   fontSize=14, leading=18, textColor=colors.white)
    est_kpi_detalle = ParagraphStyle("kpiD", parent=base, fontSize=7,
                                     leading=9, textColor=gris)
    est_seccion = ParagraphStyle("sec", parent=base, fontName="Helvetica-Bold",
                                 fontSize=11, leading=14, textColor=oscuro,
                                 spaceBefore=10, spaceAfter=4, keepWithNext=1)
    est_grafico = ParagraphStyle("gra", parent=base, fontName="Helvetica-Bold",
                                 fontSize=9, leading=12, textColor=oscuro)
    est_celda = ParagraphStyle("cel", parent=base, fontSize=8, leading=10)
    est_celda_der = ParagraphStyle("celd", parent=est_celda, alignment=TA_RIGHT)
    est_enc = ParagraphStyle("enc", parent=est_celda, fontName="Helvetica-Bold",
                             textColor=colors.white)
    est_enc_der = ParagraphStyle("encd", parent=est_enc, alignment=TA_RIGHT)

    elementos = []

    # ---------- Tarjetas de KPI ----------
    kpis = reporte["kpis"][:6]
    celdas = [
        [
            Paragraph(escape(kpi["titulo"].upper()), est_kpi_titulo),
            Paragraph(
                escape(bi_formatear(kpi["valor"], kpi.get("formato"))),
                est_kpi_valor,
            ),
            Paragraph(escape(kpi.get("detalle", "")), est_kpi_detalle),
        ]
        for kpi in kpis
    ]
    tarjetas = Table([celdas], colWidths=[ancho_total / len(celdas)] * len(celdas))
    tarjetas.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), oscuro),
        ("LINEABOVE", (0, 0), (-1, 0), 3, rojo),
        ("LINEAFTER", (0, 0), (-2, -1), 3, colors.white),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 9),
        ("RIGHTPADDING", (0, 0), (-1, -1), 6),
        ("TOPPADDING", (0, 0), (-1, -1), 8),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
    ]))
    elementos += [tarjetas, Spacer(1, 6 * mm)]

    # ---------- Gráficos (2 por fila) ----------
    graficos = reporte.get("graficos", [])
    ancho_col = (ancho_total - 8) / 2
    for inicio_fila in range(0, len(graficos), 2):
        grupo = graficos[inicio_fila:inicio_fila + 2]
        fila = [
            [
                Paragraph(escape(grafico["titulo"]), est_grafico),
                bi_dibujo_grafico(grafico, ancho_col - 14, 175),
            ]
            for grafico in grupo
        ]
        while len(fila) < 2:
            fila.append("")

        tabla_graficos = Table(
            [[fila[0], "", fila[1]]],
            colWidths=[ancho_col, 8, ancho_col],
        )
        estilo = [
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), 6),
            ("RIGHTPADDING", (0, 0), (-1, -1), 6),
            ("TOPPADDING", (0, 0), (-1, -1), 6),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ("LEFTPADDING", (1, 0), (1, 0), 0),
            ("RIGHTPADDING", (1, 0), (1, 0), 0),
            ("BOX", (0, 0), (0, 0), 0.6, borde),
        ]
        if len(grupo) == 2:
            estilo.append(("BOX", (2, 0), (2, 0), 0.6, borde))
        tabla_graficos.setStyle(TableStyle(estilo))
        elementos += [tabla_graficos, Spacer(1, 4 * mm)]

    # ---------- Tablas ----------
    for tabla in reporte.get("tablas", []):
        encabezados = tabla["encabezados"]
        formatos = tabla.get("formatos") or [None] * len(encabezados)
        derecha = [formato in BI_NUMERICOS for formato in formatos]

        datos_tabla = [[
            Paragraph(escape(enc), est_enc_der if derecha[i] else est_enc)
            for i, enc in enumerate(encabezados)
        ]]
        for fila in tabla["filas"]:
            datos_tabla.append([
                Paragraph(
                    escape(bi_formatear(valor, formatos[i])),
                    est_celda_der if derecha[i] else est_celda,
                )
                for i, valor in enumerate(fila)
            ])

        ancho_tabla = min(ancho_total, 150 * len(encabezados))
        pesos = [
            1.8 if i == 0 and not derecha[0] else 1
            for i in range(len(encabezados))
        ]
        anchos = [ancho_tabla * peso / sum(pesos) for peso in pesos]

        objeto = Table(datos_tabla, colWidths=anchos, repeatRows=1, hAlign="LEFT")
        objeto.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), rojo),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1),
             [colors.white, colors.HexColor("#F3F4F6")]),
            ("LINEBELOW", (0, 0), (-1, -1), 0.4, borde),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("LEFTPADDING", (0, 0), (-1, -1), 6),
            ("RIGHTPADDING", (0, 0), (-1, -1), 6),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ]))
        elementos += [
            Paragraph(escape(tabla["titulo"].upper()), est_seccion),
            objeto,
            Spacer(1, 3 * mm),
        ]

    marco = bi_marco_pdf(reporte["titulo"], reporte["periodo"])
    documento.build(elementos, onFirstPage=marco, onLaterPages=marco)
    return respuesta


# ------------------------------------------------------------
# 10.8 Despachador de exportaciones
# ------------------------------------------------------------

def bi_responder_exportacion(request, tipo, formato):
    reportes = {
        "ventas": (bi_reporte_ventas, "resumen_ventas"),
        "productos": (bi_reporte_productos, "productos_mas_vendidos"),
        "dinero": (bi_reporte_dinero, "dinero_recaudado"),
    }
    if tipo not in reportes:
        return HttpResponse("Reporte desconocido", status=404)
    if formato not in ("excel", "pdf"):
        return HttpResponse("Formato desconocido", status=404)

    if formato == "excel" and not OPENPYXL_DISPONIBLE:
        return HttpResponse(
            "Falta instalar openpyxl (pip install openpyxl).", status=503
        )
    if formato == "pdf" and not REPORTLAB_DISPONIBLE:
        return HttpResponse(
            "Falta instalar reportlab (pip install reportlab).", status=503
        )

    inicio, fin = bi_fechas(request)
    constructor, nombre = reportes[tipo]
    reporte = constructor(inicio, fin)
    archivo = f"{nombre}_{inicio:%Y%m%d}_{fin:%Y%m%d}"

    if formato == "excel":
        return bi_exportar_excel(reporte, f"{archivo}.xlsx")
    return bi_exportar_pdf(reporte, f"{archivo}.pdf")


# ------------------------------------------------------------
# 10.9 Vistas web de reportes
# ------------------------------------------------------------

@login_required
def centro_reportes(request):
    inicio, fin = bi_fechas(request)
    indicadores, _ = bi_resumen(inicio, fin)

    contexto = {
        **bi_contexto(inicio, fin),
        **indicadores,
        "indicadores": indicadores,
    }
    return render(request, "panel/reportes/centroReportes.html", contexto)


@login_required
def resumen_ventas(request):
    inicio, fin = bi_fechas(request)
    indicadores, ventas_diarias = bi_resumen(inicio, fin)
    metricas = bi_metricas_adicionales(inicio, fin, indicadores, ventas_diarias)

    contexto = {
        **bi_contexto(inicio, fin),
        **indicadores,
        "indicadores": indicadores,
        "ventas_diarias": ventas_diarias,
        **metricas,
    }
    return render(request, "panel/reportes/resumenVentas.html", contexto)


@login_required
def productos_mas_vendidos(request):
    inicio, fin = bi_fechas(request)
    productos = bi_top_productos(inicio, fin)
    total_unidades, total_ingresos = bi_totales_productos(productos)

    contexto = {
        **bi_contexto(inicio, fin),
        "productos": productos,
        "total_unidades": total_unidades,
        "total_ingresos": total_ingresos,
    }
    return render(request, "panel/reportes/productosMasVendidos.html", contexto)


@login_required
def dinero_recaudado(request):
    inicio, fin = bi_fechas(request)
    contexto = {**bi_contexto(inicio, fin), **bi_datos_dinero(inicio, fin)}
    return render(request, "panel/reportes/dineroRecaudado.html", contexto)


# ------------------------------------------------------------
# 10.10 Vistas de exportación
# ------------------------------------------------------------

@login_required
def exportar_resumen_excel(request):
    return bi_responder_exportacion(request, "ventas", "excel")


@login_required
def exportar_resumen_pdf(request):
    return bi_responder_exportacion(request, "ventas", "pdf")


@login_required
def exportar_productos_excel(request):
    return bi_responder_exportacion(request, "productos", "excel")


@login_required
def exportar_productos_pdf(request):
    return bi_responder_exportacion(request, "productos", "pdf")


@login_required
def exportar_dinero_excel(request):
    return bi_responder_exportacion(request, "dinero", "excel")


@login_required
def exportar_dinero_pdf(request):
    return bi_responder_exportacion(request, "dinero", "pdf")
# ------------------------------------------------------------
# VENTAS
# ------------------------------------------------------------
 
def reportes(request):
    return render(request, 'panel/reportes/reportes.html')
 
 
def reporte_resumen_ventas(request):
    hoy = timezone.localdate()
 
    rango = request.GET.get('rango', '').strip()
    fecha_desde_txt = request.GET.get('fecha_desde', '').strip()
    fecha_hasta_txt = request.GET.get('fecha_hasta', '').strip()
 
    atajo = _rango_por_atajo(rango, hoy)
 
    if atajo:
        desde, hasta = atajo
    else:
        rango = ''
        desde = (
            _fecha_o_none(fecha_desde_txt)
            if fecha_desde_txt else hoy - timedelta(days=6)
        )
        hasta = _fecha_o_none(fecha_hasta_txt) if fecha_hasta_txt else hoy
 
    filtros_validos = True
 
    if desde is None:
        messages.error(request, 'La fecha inicial no es válida.')
        filtros_validos = False
 
    if hasta is None:
        messages.error(request, 'La fecha final no es válida.')
        filtros_validos = False
 
    if desde and hasta:
        if desde > hasta:
            messages.error(
                request,
                'La fecha inicial no puede ser posterior a la fecha final.'
            )
            filtros_validos = False
 
        elif (hasta - desde).days + 1 > MAX_DIAS_REPORTE:
            messages.error(
                request,
                f'El período no puede superar los {MAX_DIAS_REPORTE} días.'
            )
            filtros_validos = False
 
    contexto = {
        'filtros_validos': filtros_validos,
        'rango_activo': rango,
        'fecha_desde': desde.isoformat() if desde else fecha_desde_txt,
        'fecha_hasta': hasta.isoformat() if hasta else fecha_hasta_txt,
    }
 
    if not filtros_validos:
        return render(request, 'panel/reportes/resumenVentas.html', contexto)
 
    exportar = request.GET.get('exportar') == 'xlsx'
 
    if exportar and Workbook is None:
        messages.error(request, MENSAJE_SIN_OPENPYXL)
        exportar = False
 
    datos = _datos_resumen_ventas(
        desde, hasta, limite_productos=50 if exportar else 5
    )
 
    if exportar:
        return _exportar_resumen_xlsx(datos, desde, hasta)
 
    contexto.update(datos)
    contexto['url_exportar'] = _url_exportar(request)
 
    return render(request, 'panel/reportes/resumenVentas.html', contexto)
 
 
def reporte_historial_ventas(request):
    """
    Historial de pedidos pagados (sin cancelados).
    Filtra por fecha inicial, fecha final y número de pedido.
    Con ?exportar=xlsx descarga todos los resultados en Excel.
    """
    pedidos = _pedidos_vendidos().select_related(
        'idcliente', 'idmediopago', 'idestadopedido'
    )
 
    fecha_desde = request.GET.get('fecha_desde', '').strip()
    fecha_hasta = request.GET.get('fecha_hasta', '').strip()
    numero_pedido = request.GET.get('pedido', '').strip()
 
    desde = _fecha_o_none(fecha_desde)
    hasta = _fecha_o_none(fecha_hasta)
 
    hay_error = False
 
    if fecha_desde and desde is None:
        messages.error(request, 'La fecha inicial no es válida.')
        hay_error = True
 
    if fecha_hasta and hasta is None:
        messages.error(request, 'La fecha final no es válida.')
        hay_error = True
 
    if desde and hasta and desde > hasta:
        messages.error(
            request,
            'La fecha inicial no puede ser posterior a la fecha final.'
        )
        hay_error = True
 
    if numero_pedido and not (
        numero_pedido.isdigit() and int(numero_pedido) > 0
    ):
        messages.error(
            request, 'El número de pedido debe ser un entero positivo.'
        )
        hay_error = True
 
    if hay_error:
        pedidos = pedidos.none()
    else:
        if desde:
            pedidos = pedidos.filter(fecha_creacion__date__gte=desde)
 
        if hasta:
            pedidos = pedidos.filter(fecha_creacion__date__lte=hasta)
 
        if numero_pedido:
            pedidos = pedidos.filter(idpedido=int(numero_pedido))
 
    pedidos = pedidos.order_by('-fecha_creacion', '-idpedido')
 
    if request.GET.get('exportar') == 'xlsx' and not hay_error:
        if Workbook is None:
            messages.error(request, MENSAJE_SIN_OPENPYXL)
        else:
            partes = []
 
            if desde:
                partes.append(f'desde {desde:%d/%m/%Y}')
 
            if hasta:
                partes.append(f'hasta {hasta:%d/%m/%Y}')
 
            if numero_pedido:
                partes.append(f'pedido #{numero_pedido}')
 
            return _exportar_historial_xlsx(
                pedidos,
                'Filtros: ' + (' · '.join(partes) or 'todos los pedidos'),
            )
 
    resumen = pedidos.aggregate(
        total_ventas=Sum('total'),
        cantidad_pedidos=Count('idpedido'),
        ticket_promedio=Avg('total'),
    )
 
    pagina = Paginator(pedidos, 15).get_page(request.GET.get('pagina'))
 
    return render(request, 'panel/reportes/historialVentas.html', {
        'pagina_pedidos': pagina,
        'total_ventas': resumen['total_ventas'] or CERO,
        'cantidad_pedidos': resumen['cantidad_pedidos'] or 0,
        'ticket_promedio': resumen['ticket_promedio'] or CERO,
        'fecha_desde': fecha_desde,
        'fecha_hasta': fecha_hasta,
        'numero_pedido': numero_pedido,
        'url_exportar': _url_exportar(request),
    })
# ------------------------------------------------------------
# HORARIOS
# ------------------------------------------------------------
class HorarioForm(forms.ModelForm):
    DIAS_SEMANA = [
        ("Lunes", "Lunes"),
        ("Martes", "Martes"),
        ("Miércoles", "Miércoles"),
        ("Jueves", "Jueves"),
        ("Viernes", "Viernes"),
        ("Sábado", "Sábado"),
        ("Domingo", "Domingo"),
    ]

    ESTADOS = [
        (1, "Abierto"),
        (0, "Cerrado"),
    ]

    diasemana = forms.ChoiceField(
        choices=DIAS_SEMANA,
        label="Día de la semana",
        widget=forms.Select(attrs={"class": "form-control"})
    )

    idnegocio = forms.ModelChoiceField(
        queryset=Negocio.objects.all(),
        label="Negocio",
        empty_label="Seleccioná un negocio",
        widget=forms.Select(attrs={"class": "form-control"})
    )

    horaapertura = forms.TimeField(
        label="Hora de apertura",
        required=False,
        input_formats=["%H:%M"],
        widget=forms.TimeInput(
            format="%H:%M",
            attrs={"type": "time", "class": "form-control"}
        )
    )

    horacierre = forms.TimeField(
        label="Hora de cierre",
        required=False,
        input_formats=["%H:%M"],
        widget=forms.TimeInput(
            format="%H:%M",
            attrs={"type": "time", "class": "form-control"}
        )
    )

    estaabierto = forms.TypedChoiceField(
        choices=ESTADOS,
        label="Estado del día",
        coerce=int,
        widget=forms.Select(attrs={"class": "form-control"})
    )

    class Meta:
        model = Horario
        fields = [
            "idnegocio",
            "diasemana",
            "horaapertura",
            "horacierre",
            "estaabierto",
        ]

def horarios_lista(request):
    horarios = Horario.objects.select_related("idnegocio").order_by(
        "idnegocio__nombre", "idhorario"
    )
    return render(
        request,
        "panel/horarios/horarioLista.html",
        {"horarios": horarios},
    )


def horarios_crear(request):
    if request.method == "POST":
        form = HorarioForm(request.POST)
        if form.is_valid():
            form.save()
            messages.success(request, "Horario creado correctamente.")
            return redirect("horarios_lista")
    else:
        form = HorarioForm()

    return render(request, "panel/horarios/horarioCrear.html", {
        "form": form,
        "titulo": "Crear horario",
        "accion": "Crear horario",
    })


def horarios_editar(request, idhorario):
    horario = get_object_or_404(Horario, pk=idhorario)

    if request.method == "POST":
        form = HorarioForm(request.POST, instance=horario)
        if form.is_valid():
            form.save()
            messages.success(request, "Horario actualizado correctamente.")
            return redirect("horarios_lista")
    else:
        form = HorarioForm(instance=horario)

    return render(request, "panel/horarios/horarioEditar.html", {
        "form": form,
        "titulo": "Editar horario",
        "accion": "Guardar cambios",
    })


def horarios_eliminar(request, idhorario):
    horario = get_object_or_404(Horario, pk=idhorario)

    if request.method == "POST":
        horario.delete()
        messages.success(request, "Horario eliminado correctamente.")
        return redirect("horarios_lista")

    return render(
        request,
        "panel/horarios/horarioEliminar.html",
        {"horario": horario},
    )
# ------------------------------------------------------------
# ZONAS ENTREGA
# ------------------------------------------------------------
def zonas_entrega_lista(request):
    zonas = ZonasEntrega.objects.all().order_by("nombre", "idzona")

    return render(
        request,
        "panel/entrega/zonaLista.html",
        {"zonas": zonas},
    )


def zonas_entrega_crear(request):
    if request.method == "POST":
        nombre = request.POST.get("nombre", "").strip()
        costoenvio = request.POST.get("costoenvio") or "0"
        idnegocio = request.POST.get("idnegocio") or None

        try:
            costoenvio = Decimal(costoenvio)

            if costoenvio < 0:
                raise InvalidOperation

            if not nombre:
                raise ValueError("El nombre es obligatorio.")

            if len(nombre) > 50:
                raise ValueError("El nombre no puede superar los 50 caracteres.")

            if idnegocio is not None:
                idnegocio = int(idnegocio)
                if idnegocio < 1:
                    raise ValueError("El ID del negocio debe ser positivo.")

            ZonasEntrega.objects.create(
                nombre=nombre,
                costoenvio=costoenvio,
                idnegocio=idnegocio,
            )

            messages.success(request, "Zona de entrega creada correctamente.")
            return redirect("zonas_entrega_lista")

        except (ValueError, InvalidOperation, TypeError):
            messages.error(
                request,
                "Verificá los datos: nombre obligatorio, costo válido e ID de negocio numérico.",
            )

    return render(request, "panel/entrega/zonaCrear.html")


def zonas_entrega_editar(request, idzona):
    zona = get_object_or_404(ZonasEntrega, pk=idzona)

    if request.method == "POST":
        nombre = request.POST.get("nombre", "").strip()
        costoenvio = request.POST.get("costoenvio") or "0"
        idnegocio = request.POST.get("idnegocio") or None

        try:
            costoenvio = Decimal(costoenvio)

            if costoenvio < 0:
                raise InvalidOperation

            if not nombre:
                raise ValueError("El nombre es obligatorio.")

            if len(nombre) > 50:
                raise ValueError("El nombre no puede superar los 50 caracteres.")

            if idnegocio is not None:
                idnegocio = int(idnegocio)
                if idnegocio < 1:
                    raise ValueError("El ID del negocio debe ser positivo.")

            zona.nombre = nombre
            zona.costoenvio = costoenvio
            zona.idnegocio = idnegocio
            zona.save()

            messages.success(request, "Zona de entrega actualizada correctamente.")
            return redirect("zonas_entrega_lista")

        except (ValueError, InvalidOperation, TypeError):
            messages.error(
                request,
                "Verificá los datos: nombre obligatorio, costo válido e ID de negocio numérico.",
            )

    return render(
        request,
        "panel/entrega/zonaEditar.html",
        {"zona": zona},
    )


def zonas_entrega_eliminar(request, idzona):
    zona = get_object_or_404(ZonasEntrega, pk=idzona)

    if request.method == "POST":
        zona.delete()
        messages.success(request, "Zona de entrega eliminada correctamente.")
        return redirect("zonas_entrega_lista")

    return render(
        request,
        "panel/entrega/zonaEliminar.html",
        {"zona": zona},
    )
# ------------------------------------------------------------
# historial de pedidos
# ------------------------------------------------------------
def historial_pedidos(request):
    """Historial general de pedidos, independientemente de su estado."""
    pedidos = Pedido.objects.select_related(
        'idcliente',
        'idmediopago',
        'idestadopedido',
    ).all()

    fecha_desde = request.GET.get('fecha_desde', '').strip()
    fecha_hasta = request.GET.get('fecha_hasta', '').strip()
    numero_pedido = request.GET.get('pedido', '').strip()
    cliente = request.GET.get('cliente', '').strip()
    estado = request.GET.get('estado', '').strip()

    desde = _fecha_o_none(fecha_desde) if fecha_desde else None
    hasta = _fecha_o_none(fecha_hasta) if fecha_hasta else None
    hay_error = False

    if fecha_desde and desde is None:
        messages.error(request, 'La fecha inicial no es válida.')
        hay_error = True

    if fecha_hasta and hasta is None:
        messages.error(request, 'La fecha final no es válida.')
        hay_error = True

    if desde and hasta and desde > hasta:
        messages.error(
            request,
            'La fecha inicial no puede ser posterior a la fecha final.'
        )
        hay_error = True

    if numero_pedido and not (
        numero_pedido.isdigit() and int(numero_pedido) > 0
    ):
        messages.error(
            request,
            'El número de pedido debe ser un entero positivo.'
        )
        hay_error = True

    if estado and not estado.isdigit():
        messages.error(request, 'El estado seleccionado no es válido.')
        hay_error = True

    if hay_error:
        pedidos = pedidos.none()
    else:
        if desde:
            pedidos = pedidos.filter(fecha_creacion__date__gte=desde)

        if hasta:
            pedidos = pedidos.filter(fecha_creacion__date__lte=hasta)

        if numero_pedido:
            pedidos = pedidos.filter(idpedido=int(numero_pedido))

        if cliente:
            pedidos = pedidos.filter(
                Q(idcliente__nombre__icontains=cliente)
                | Q(idcliente__apellido__icontains=cliente)
            )

        if estado:
            pedidos = pedidos.filter(idestadopedido__idestado=int(estado))

    pedidos = pedidos.order_by('-fecha_creacion', '-idpedido')
    estados = EstadoPedido.objects.all().order_by('descripcion')

    pagina = Paginator(pedidos, 15).get_page(
        request.GET.get('pagina')
    )

    return render(request, 'panel/ventas/historialPedidos.html', {
        'pagina_pedidos': pagina,
        'estados': estados,
        'fecha_desde': fecha_desde,
        'fecha_hasta': fecha_hasta,
        'numero_pedido': numero_pedido,
        'cliente': cliente,
        'estado_seleccionado': estado,
    })


def detalle_historial_pedido(request, idpedido):
    """Detalle completo de un pedido del historial."""
    pedido = get_object_or_404(
        Pedido.objects.select_related(
            'idcliente',
            'idmediopago',
            'idestadopedido',
        ),
        idpedido=idpedido,
    )

    detalles = DetallePedido.objects.filter(
        idpedido=pedido
    ).select_related('idproducto')

    for detalle in detalles:
        detalle.subtotal_item = detalle.preciounitario * detalle.cantidad

    return render(request, 'panel/ventas/detalleHistorialPedido.html', {
        'pedido': pedido,
        'detalles': detalles,
    })
# ============================================================
# 11. PANEL - ABRIR / CERRAR EL NEGOCIO
# ============================================================

# ------------------------------------------------------------
# 11.1 Constantes
# ------------------------------------------------------------

# Estados fijos de EstadoNegocio (se crean solos la primera vez que se usan)
LOCAL_AUTOMATICO = "Automático"
LOCAL_ABIERTO_MANUAL = "Abierto manual"
LOCAL_CERRADO_MANUAL = "Cerrado manual"

# Prefijo con el que se registran los cambios en UltimaModificacion
LOCAL_PREFIJO_LOG = "Estado del local:"

# Si no hay horarios cargados, ¿se considera abierto?
# True evita bloquear pedidos mientras no se carguen los horarios.
LOCAL_SIN_HORARIOS_ABIERTO = True

LOCAL_DIAS = [
    "lunes", "martes", "miercoles", "jueves",
    "viernes", "sabado", "domingo",
]
LOCAL_DIAS_TXT = [
    "lunes", "martes", "miércoles", "jueves",
    "viernes", "sábado", "domingo",
]

LOCAL_TEXTO_LOG = {
    LOCAL_AUTOMATICO: "volvió al modo automático (según horarios)",
    LOCAL_ABIERTO_MANUAL: "abierto manualmente por hoy",
    LOCAL_CERRADO_MANUAL: "cerrado manualmente por hoy",
}


# ------------------------------------------------------------
# 11.2 Utilidades
# ------------------------------------------------------------

def local_normalizar(texto):
    texto = unicodedata.normalize("NFD", texto or "")
    sin_tildes = "".join(c for c in texto if unicodedata.category(c) != "Mn")
    return sin_tildes.strip().lower()


LOCAL_INDICE_DIA = {nombre: i for i, nombre in enumerate(LOCAL_DIAS)}
LOCAL_MODOS = {
    local_normalizar(modo): modo
    for modo in (LOCAL_AUTOMATICO, LOCAL_ABIERTO_MANUAL, LOCAL_CERRADO_MANUAL)
}


def local_ahora():
    ahora = timezone.now()
    return timezone.localtime(ahora) if timezone.is_aware(ahora) else ahora


def local_a_local(momento):
    return timezone.localtime(momento) if timezone.is_aware(momento) else momento


def local_nombre_empleado(empleado):
    if empleado is None:
        return ""
    return f"{empleado.nombre} {getattr(empleado, 'apellido', '')}".strip()


# ------------------------------------------------------------
# 11.3 Horarios automáticos
# ------------------------------------------------------------

def local_franjas(horarios):
    """[(día 0-6, apertura, cierre)] solo de los días marcados como abiertos."""
    franjas = []
    for h in horarios:
        if not (h.estaabierto and h.horaapertura and h.horacierre):
            continue
        dia = LOCAL_INDICE_DIA.get(local_normalizar(h.diasemana))
        if dia is not None:
            franjas.append((dia, h.horaapertura, h.horacierre))
    return franjas


def local_abierto_por_horario(franjas, ahora):
    dia, hora = ahora.weekday(), ahora.time()
    for d, apertura, cierre in franjas:
        if apertura == cierre:                       # 24 horas
            if d == dia:
                return True
        elif apertura < cierre:                      # mismo día
            if d == dia and apertura <= hora < cierre:
                return True
        else:                                        # cruza la medianoche
            if d == dia and hora >= apertura:
                return True
            if d == (dia - 1) % 7 and hora < cierre:
                return True
    return False


def local_proxima_apertura(franjas, desde):
    candidatas = []
    for desplazamiento in range(8):
        fecha = desde.date() + timedelta(days=desplazamiento)
        for d, apertura, _ in franjas:
            if d == fecha.weekday():
                momento = datetime.combine(fecha, apertura, tzinfo=desde.tzinfo)
                if momento > desde:
                    candidatas.append(momento)
    return min(candidatas) if candidatas else None


def local_texto_apertura(momento, ahora):
    dias = (momento.date() - ahora.date()).days
    if dias == 0:
        cuando = "hoy"
    elif dias == 1:
        cuando = "mañana"
    else:
        cuando = f"el {LOCAL_DIAS_TXT[momento.weekday()]}"
    return f"{cuando} a las {momento:%H:%M}"


# ------------------------------------------------------------
# 11.4 Estado actual del local
# ------------------------------------------------------------

def local_estado_por_defecto():
    return {
        "abierto": True,
        "manual": False,
        "mensaje": "",
        "ultimo_cambio": None,
        "ultimo_por": "",
    }


def local_estado():
    """
    Estado actual del local (se usa en el panel y, más adelante,
    para bloquear carrito y checkout). Devuelve un diccionario:
    abierto, manual (hay un cierre/apertura manual vigente HOY),
    mensaje, ultimo_cambio y ultimo_por.
    """
    try:
        return local_calcular_estado()
    except Exception:
        logger.exception("No se pudo calcular el estado del local")
        return local_estado_por_defecto()


def local_calcular_estado():
    ahora = local_ahora()
    negocio = Negocio.objects.select_related("idestado").first()
    if negocio is None:
        return local_estado_por_defecto()

    descripcion = negocio.idestado.descripcion if negocio.idestado_id else ""
    modo = LOCAL_MODOS.get(local_normalizar(descripcion), LOCAL_AUTOMATICO)

    registro = (
        UltimaModificacion.objects.select_related("idempleado")
        .filter(
            idnegocio=negocio.idnegocio,
            descripcioncambio__startswith=LOCAL_PREFIJO_LOG,
        )
        .order_by("-fecha")
        .first()
    )

    # El modo manual solo vale el día en que se activó
    manual_vigente = (
        modo != LOCAL_AUTOMATICO
        and registro is not None
        and local_a_local(registro.fecha).date() == ahora.date()
    )

    franjas = local_franjas(Horario.objects.filter(idnegocio=negocio))

    if manual_vigente:
        abierto = modo == LOCAL_ABIERTO_MANUAL
    elif franjas:
        abierto = local_abierto_por_horario(franjas, ahora)
    else:
        abierto = LOCAL_SIN_HORARIOS_ABIERTO

    if abierto:
        mensaje = "Estamos abiertos."
    else:
        if manual_vigente:
            medianoche = datetime.combine(
                ahora.date() + timedelta(days=1),
                datetime.min.time(),
                tzinfo=ahora.tzinfo,
            )
            proxima = local_proxima_apertura(
                franjas, medianoche - timedelta(seconds=1)
            )
            mensaje = "Hoy el local permanece cerrado."
        else:
            proxima = local_proxima_apertura(franjas, ahora)
            mensaje = "En este momento estamos cerrados."
        if proxima:
            mensaje += f" Volvemos a abrir {local_texto_apertura(proxima, ahora)}."

    return {
        "abierto": abierto,
        "manual": manual_vigente,
        "mensaje": mensaje,
        "ultimo_cambio": local_a_local(registro.fecha) if registro else None,
        "ultimo_por": local_nombre_empleado(registro.idempleado) if registro else "",
    }


# ------------------------------------------------------------
# 11.5 Cambio de modo (abrir / cerrar / automático)
# ------------------------------------------------------------

def local_estado_fijo(descripcion, id_negocio, ahora):
    estado = EstadoNegocio.objects.filter(descripcion=descripcion).first()
    if estado:
        return estado
    modificacion = UltimaModificacion.objects.create(
        fecha=ahora,
        idempleado=None,
        descripcioncambio=f"Estado inicial: {descripcion}",
        idnegocio=id_negocio,
    )
    return EstadoNegocio.objects.create(
        idmodificacion=modificacion, descripcion=descripcion
    )


def local_aplicar_modo(modo, empleado=None):
    """Guarda el nuevo modo y deja registro de quién y cuándo. Devuelve el Negocio."""
    with transaction.atomic():
        negocio = Negocio.objects.select_for_update().first()
        if negocio is None:
            return None

        ahora = local_ahora()
        estado = local_estado_fijo(modo, negocio.idnegocio, ahora)

        UltimaModificacion.objects.create(
            fecha=ahora,
            idempleado=empleado,
            descripcioncambio=f"{LOCAL_PREFIJO_LOG} {LOCAL_TEXTO_LOG[modo]}",
            idnegocio=negocio.idnegocio,
        )

        negocio.idestado = estado
        negocio.idempleado = empleado
        negocio.save(update_fields=["idestado", "idempleado"])
    return negocio


# ------------------------------------------------------------
# 11.6 Vistas del panel
# ------------------------------------------------------------

@login_required
def estado_negocio_panel(request):
    estado = local_estado()
    negocio = Negocio.objects.first()

    cambios = []
    total_aperturas = total_cierres = total_automatico = 0

    if negocio:
        base = UltimaModificacion.objects.filter(
            idnegocio=negocio.idnegocio,
            descripcioncambio__startswith=LOCAL_PREFIJO_LOG,
        )

        def contar(modo):
            return base.filter(
                descripcioncambio=f"{LOCAL_PREFIJO_LOG} {LOCAL_TEXTO_LOG[modo]}"
            ).count()

        total_aperturas = contar(LOCAL_ABIERTO_MANUAL)
        total_cierres = contar(LOCAL_CERRADO_MANUAL)
        total_automatico = contar(LOCAL_AUTOMATICO)

        modos_por_texto = {texto: modo for modo, texto in LOCAL_TEXTO_LOG.items()}
        acciones = {
            LOCAL_ABIERTO_MANUAL: ("Abrió el local", "abrio"),
            LOCAL_CERRADO_MANUAL: ("Cerró el local", "cerro"),
            LOCAL_AUTOMATICO: ("Volvió a modo automático", "auto"),
        }

        registros = (
            base.select_related("idempleado")
            .order_by("-fecha")[:100]
        )

        for posicion, registro in enumerate(registros):
            texto = (registro.descripcioncambio or "")[len(LOCAL_PREFIJO_LOG):].strip()
            modo = modos_por_texto.get(texto)
            accion, clase = acciones.get(modo, (texto or "Cambio", "auto"))

            cambios.append({
                "fecha": local_a_local(registro.fecha),
                "accion": accion,
                "clase": clase,
                "quien": (
                    local_nombre_empleado(registro.idempleado)
                    or "Sin usuario registrado"
                ),
                "vigente": posicion == 0 and estado["manual"],
            })

    contexto = {
        "estado_local": estado,
        "cambios": cambios,
        "total_aperturas": total_aperturas,
        "total_cierres": total_cierres,
        "total_automatico": total_automatico,
    }
    return render(request, "panel/negocio/interruptorNegocio.html", contexto)


@login_required
@require_POST
def cambiar_estado_local(request):
    destino = request.POST.get("next", "")
    if not url_has_allowed_host_and_scheme(
        destino,
        allowed_hosts={request.get_host()},
        require_https=request.is_secure(),
    ):
        destino = "/"

    modos = {
        "abrir": (
            LOCAL_ABIERTO_MANUAL,
            "El local quedó abierto por hoy. Mañana vuelve a regirse por los horarios.",
        ),
        "cerrar": (
            LOCAL_CERRADO_MANUAL,
            "El local quedó cerrado por hoy. Mañana abre normalmente según los horarios.",
        ),
        "automatico": (
            LOCAL_AUTOMATICO,
            "El local volvió al modo automático según los horarios.",
        ),
    }

    accion = request.POST.get("accion")
    if accion not in modos:
        messages.error(request, "Acción no válida.")
        return redirect(destino)

    modo, texto = modos[accion]
    empleado = request.user if isinstance(request.user, Empleado) else None
    negocio = local_aplicar_modo(modo, empleado)

    if negocio is None:
        messages.error(request, "No hay un negocio configurado.")
    else:
        messages.success(request, texto)

    return redirect(destino)